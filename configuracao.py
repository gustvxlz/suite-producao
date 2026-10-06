"""Configuração persistente e editável da Suite de Produção.

As preferências ficam fora do código, em um JSON do usuário. Assim upgrades do
programa não apagam grupos, exclusões ou regras personalizadas.
"""

from __future__ import annotations

import copy
import json
import os
import shutil
import threading
import queue
import unicodedata
import re
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext

import openpyxl

APP_NAME = "SuiteProducao"
CONFIG_VERSION = 3


def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = texto.upper()
    texto = re.sub(r"[^A-Z0-9 ]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


DEFAULT_RULES = [
    {"nome": "Mesa Salinas/Moscou", "acao": "renomear", "todos": ["MESA"], "algum": ["SALINAS", "MOSCOU"], "nenhum": [], "destino": "MESA HAPPY HOUR", "nota": "renomeado: Mesa Salinas/Moscou -> Mesa Happy Hour", "ativo": True},
    {"nome": "Mesa Santiago", "acao": "renomear", "todos": ["MESA", "SANTIAGO"], "algum": [], "nenhum": [], "destino": "MESA SANTIAGO", "nota": "padronizado: Mesa Santiago", "ativo": True},
    {"nome": "Trama Mesa", "acao": "renomear", "todos": ["TRAMA", "MESA"], "algum": [], "nenhum": [], "destino": "MESA NAVEGANTES", "nota": "renomeado: Trama Mesa -> Mesa Nave", "ativo": True},
    {"nome": "Mesa Happy Hour", "acao": "renomear", "todos": ["MESA", "HAPPY", "HOUR"], "algum": [], "nenhum": [], "destino": "MESA HAPPY HOUR", "nota": "padronizado: Mesa Happy Hour", "ativo": True},
    {"nome": "Namo Santiago", "acao": "renomear", "todos": ["NAMO", "SANTIAGO"], "algum": [], "nenhum": [], "destino": "NAMO SANTIAGO", "nota": "padronizado: Namo Santiago", "ativo": True},
    {"nome": "Trama Namo", "acao": "renomear", "todos": ["TRAMA", "NAMO"], "algum": [], "nenhum": [], "destino": "NAMO NAVE", "nota": "renomeado: Trama Namo -> Namo Nave", "ativo": True},
    {"nome": "Poltrona Santiago", "acao": "renomear", "todos": ["SANTIAGO"], "algum": ["POLT", "POLTRONA"], "nenhum": [], "destino": "POLT SANTIAGO", "nota": "padronizado: Polt Santiago", "ativo": True},
    {"nome": "Trama Poltrona", "acao": "renomear", "todos": ["TRAMA"], "algum": ["POLT", "POLTRONA"], "nenhum": [], "destino": "POLT NAVE", "nota": "renomeado: Trama Polt -> Polt Nave", "ativo": True},
    {"nome": "Gota Tramado", "acao": "renomear", "todos": ["GOTA", "TRAMADO"], "algum": [], "nenhum": [], "destino": "BALANÇO NINHO", "nota": "renomeado: Gota Tramado -> Balanço Ninho", "ativo": True},
    {"nome": "Ninho Corda/Trico", "acao": "renomear", "todos": ["NINHO"], "algum": ["CORDA", "TRICO"], "nenhum": [], "destino": "NINHO CORDA/TRICO", "nota": "padronizado automaticamente: Ninho Corda/Trico", "ativo": True},
    {"nome": "Ninho Fibra/Arg", "acao": "renomear", "todos": ["NINHO"], "algum": ["FIBRA", "ARG"], "nenhum": [], "destino": "BALANÇO NINHO", "nota": "renomeado: Ninho Fibra/Arg -> Balanço Ninho", "ativo": True},
    {"nome": "Porto", "acao": "renomear", "todos": ["PORTO"], "algum": [], "nenhum": ["MESA", "POLT", "POLTRONA"], "destino": "DESCANSO PORTO", "nota": "renomeado: Porto -> Descanso Porto", "ativo": True},
    {"nome": "Vigor", "acao": "renomear", "todos": ["VIGOR"], "algum": [], "nenhum": ["BANQ"], "destino": "CADE VIGOR", "nota": "renomeado: Vigor Náutica -> Cade Vigor", "ativo": True},
    {"nome": "Veneza", "acao": "renomear", "todos": ["VENEZA"], "algum": [], "nenhum": [], "destino": "CADE VENEZA", "nota": "padronizado: Cade Veneza", "ativo": True},
    {"nome": "Happy Hour cadeira", "acao": "renomear", "todos": ["HAPPY", "HOUR"], "algum": [], "nenhum": ["MESA"], "destino": "CADE HAPPY HOUR", "nota": "padronizado: Cade Happy Hour", "ativo": True},
    {"nome": "Forts", "acao": "renomear", "todos": ["FORTS"], "algum": [], "nenhum": [], "destino": "BANQ FORTS", "nota": "padronizado: Banq Forts", "ativo": True},
]

DEFAULT_CONFIG = {
    "version": CONFIG_VERSION,
    # Cadastre seus grupos na interface. Nenhum funcionário real é distribuído.
    "grupos": {},
    "exclusoes_produtos": [
        {"texto": "REFORMA", "modo": "contem"},
        {"texto": "BORDA", "modo": "contem"},
        {"texto": "ENCOSTO", "modo": "contem"},
    ],
    "regras_produtos": DEFAULT_RULES,
    "catalogo_regras": {
        "planilhas": [],
        "produtos": [],
    },
    "comentarios_funcionario": [
        "FALTA", "ATESTADO", "FOLGA", "FERIAS", "ATRASO", "SAIDA", "DISPENSA", "LICENCA",
    ],
    "matching": {"aceitar": 85, "revisar": 60, "funcionario": 65},
    "preferencias": {
        "somar_existente": True,
        "lembrar_arquivos": True,
        "confirmar_salvamento": True,
    },
    "historico": {
        "extrator_arquivo": "",
        "cadeiras_arquivo": "",
        "mesas_arquivo": "",
        "individual_arquivo": "",
        "pasta_saida": "",
    },
}


def _config_dir() -> Path:
    override = os.environ.get("SUITE_PRODUCAO_CONFIG_DIR")
    if override:
        return Path(override)
    if os.name == "nt":
        base = os.environ.get("APPDATA") or str(Path.home())
        return Path(base) / APP_NAME
    base = os.environ.get("XDG_CONFIG_HOME")
    if base:
        return Path(base) / APP_NAME
    return Path.home() / ".config" / APP_NAME


CONFIG_DIR = _config_dir()
CONFIG_PATH = CONFIG_DIR / "config.json"


def _merge_defaults(data, defaults):
    """Mescla chaves novas sem apagar customizações do usuário."""
    if not isinstance(data, dict):
        return copy.deepcopy(defaults)
    result = copy.deepcopy(data)
    for key, value in defaults.items():
        if key not in result:
            result[key] = copy.deepcopy(value)
        elif isinstance(value, dict):
            result[key] = _merge_defaults(result[key], value)
    return result


class ConfigStore:
    def __init__(self, path: Path = CONFIG_PATH):
        self.path = Path(path)
        self._lock = threading.RLock()
        self._data = copy.deepcopy(DEFAULT_CONFIG)
        self.load()

    def load(self):
        with self._lock:
            try:
                if self.path.exists():
                    with self.path.open("r", encoding="utf-8") as f:
                        loaded = json.load(f)
                    self._data = _merge_defaults(loaded, DEFAULT_CONFIG)
                else:
                    self._data = copy.deepcopy(DEFAULT_CONFIG)
                    self.save()
            except Exception:
                # Arquivo corrompido: preserva para recuperação e volta aos padrões.
                try:
                    broken = self.path.with_suffix(".corrompido.json")
                    if self.path.exists():
                        shutil.copy2(self.path, broken)
                except Exception:
                    pass
                self._data = copy.deepcopy(DEFAULT_CONFIG)
                self.save()
            self._data["version"] = CONFIG_VERSION
            return copy.deepcopy(self._data)

    def save(self):
        with self._lock:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            temp = self.path.with_suffix(".tmp")
            with temp.open("w", encoding="utf-8") as f:
                json.dump(self._data, f, ensure_ascii=False, indent=2)
                f.flush()
                try:
                    os.fsync(f.fileno())
                except OSError:
                    pass
            os.replace(temp, self.path)

    def data(self):
        with self._lock:
            return copy.deepcopy(self._data)

    def get(self, *keys, default=None):
        with self._lock:
            cur = self._data
            for key in keys:
                if not isinstance(cur, dict) or key not in cur:
                    return copy.deepcopy(default)
                cur = cur[key]
            return copy.deepcopy(cur)

    def set(self, *keys, value):
        if not keys:
            raise ValueError("Informe ao menos uma chave")
        with self._lock:
            cur = self._data
            for key in keys[:-1]:
                nxt = cur.get(key)
                if not isinstance(nxt, dict):
                    nxt = {}
                    cur[key] = nxt
                cur = nxt
            cur[keys[-1]] = copy.deepcopy(value)
            self.save()

    def replace_section(self, key, value):
        with self._lock:
            self._data[key] = copy.deepcopy(value)
            self.save()

    def reset(self):
        with self._lock:
            self._data = copy.deepcopy(DEFAULT_CONFIG)
            self.save()

    def export_to(self, path):
        with self._lock:
            shutil.copy2(self.path, path)

    def import_from(self, path):
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            raise ValueError("O arquivo de configuração não é válido.")
        with self._lock:
            self._data = _merge_defaults(data, DEFAULT_CONFIG)
            self._data["version"] = CONFIG_VERSION
            self.save()


STORE = ConfigStore()


def get_config():
    return STORE.data()


def salvar_historico(chave: str, valor: str):
    if STORE.get("preferencias", "lembrar_arquivos", default=True):
        STORE.set("historico", chave, value=str(valor or ""))


def _produto_valido(valor):
    if not isinstance(valor, str):
        return None
    nome = valor.strip()
    if not nome:
        return None
    norm = _normalizar(nome)
    if norm in {"PRODUTO", "PRODUTOS", "ESPECIFICACAO", "TOTAIS", "TOTAL", "N", "ORD"}:
        return None
    return nome


def extrair_catalogo_planilha(caminho: str):
    """Lê nomes de produtos de layouts usados pela Suite.

    Retorna itens com nome/arquivo/aba. A função é propositalmente tolerante:
    reconhece tanto a planilha CADEIRAS/MESAS (ESPECIFICAÇÃO), quanto a mensal
    e a planilha do Extrator (PRODUTOS).
    """
    caminho = os.path.abspath(os.path.expanduser(str(caminho or "")))
    if not os.path.isfile(caminho):
        raise FileNotFoundError(f"Planilha não encontrada: {caminho}")

    keep_vba = caminho.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=False, keep_vba=keep_vba)
    itens = []
    vistos = set()
    try:
        for ws in wb.worksheets:
            if "TOTAL" in ws.title.upper():
                continue

            produto_col = None
            header_row = None
            # Procura os cabeçalhos conhecidos só no começo da planilha para ser rápido.
            max_scan_row = min(ws.max_row or 1, 20)
            max_scan_col = min(ws.max_column or 1, 12)
            for row in ws.iter_rows(min_row=1, max_row=max_scan_row, min_col=1, max_col=max_scan_col):
                achou = False
                for cell in row:
                    norm = _normalizar(cell.value)
                    if norm in {"ESPECIFICACAO", "PRODUTOS"}:
                        produto_col = cell.column
                        header_row = cell.row
                        achou = True
                        break
                if achou:
                    break

            # Layout mensal: B1 contém "DIA" e os produtos ficam na coluna A.
            if produto_col is None:
                b1 = ws.cell(row=1, column=2).value
                if isinstance(b1, str) and "DIA" in _normalizar(b1):
                    produto_col = 1
                    header_row = 2

            if produto_col is None:
                continue

            for r in range((header_row or 1) + 1, (ws.max_row or 0) + 1):
                nome = _produto_valido(ws.cell(row=r, column=produto_col).value)
                if not nome:
                    continue
                chave = _normalizar(nome)
                # Deduplica pelo nome. Para atribuição, o que importa é o texto real do produto.
                if chave in vistos:
                    continue
                vistos.add(chave)
                itens.append({
                    "nome": nome,
                    "arquivo": os.path.basename(caminho),
                    "caminho": caminho,
                    "aba": ws.title,
                })
    finally:
        wb.close()
    return itens


def catalogo_regras():
    sec = STORE.get("catalogo_regras", default={}) or {}
    produtos = sec.get("produtos", []) or []
    # Compatibilidade caso alguma versão tenha salvo apenas strings.
    return [p if isinstance(p, dict) else {"nome": str(p), "arquivo": "", "aba": ""} for p in produtos]


def salvar_catalogo_regras(planilhas, produtos):
    planilhas_limpas = []
    vistos = set()
    for path in planilhas or []:
        path = os.path.abspath(os.path.expanduser(str(path)))
        key = os.path.normcase(path)
        if path and key not in vistos:
            vistos.add(key)
            planilhas_limpas.append(path)
    STORE.replace_section("catalogo_regras", {"planilhas": planilhas_limpas, "produtos": produtos or []})


def thresholds():
    m = STORE.get("matching", default={}) or {}
    return int(m.get("aceitar", 85)), int(m.get("revisar", 60)), int(m.get("funcionario", 65))


def _exclusao_bate(nome: str, item: dict) -> bool:
    alvo = _normalizar(nome)
    texto = _normalizar(item.get("texto", ""))
    if not texto:
        return False
    if item.get("modo") == "exato":
        return alvo == texto
    return texto in alvo


def exclusao_permanente(nome: str):
    """Retorna (True, motivo) quando o produto foi salvo para nunca entrar."""
    for item in STORE.get("exclusoes_produtos", default=[]) or []:
        if isinstance(item, str):
            item = {"texto": item, "modo": "contem"}
        if _exclusao_bate(nome, item):
            return True, f"exclusão permanente: {item.get('texto', '')}"
    return False, None


def adicionar_exclusao(nome: str, modo="exato"):
    nome = str(nome or "").strip()
    if not nome:
        return False
    itens = STORE.get("exclusoes_produtos", default=[]) or []
    norm = _normalizar(nome)
    for item in itens:
        if isinstance(item, str):
            texto, item_modo = item, "contem"
        else:
            texto, item_modo = item.get("texto", ""), item.get("modo", "contem")
        if _normalizar(texto) == norm and item_modo == modo:
            return False
    itens.append({"texto": nome, "modo": modo})
    STORE.replace_section("exclusoes_produtos", itens)
    return True


def aplicar_regras_produto(nome_bruto: str):
    """Aplica exclusões permanentes, atribuições simples e regras avançadas.

    Retorna (nome_para_busca, nota). None significa excluir; a constante textual
    __SEM_CATALOGO__ significa deixar pendente para decisão manual.
    """
    excluido, motivo = exclusao_permanente(nome_bruto)
    if excluido:
        return None, motivo

    t = _normalizar(nome_bruto)
    for regra in STORE.get("regras_produtos", default=[]) or []:
        if not regra.get("ativo", True):
            continue

        # Regra simples criada pelo editor "quando chegar X -> virar Y".
        if regra.get("tipo") == "mapeamento":
            origem = _normalizar(regra.get("origem", ""))
            if not origem:
                continue
            modo = regra.get("modo_origem", "contem")
            bate = (t == origem) if modo == "exato" else (origem in t)
            if not bate:
                continue
            acao = regra.get("acao", "renomear")
            nota = regra.get("nota") or f"atribuição salva: {regra.get('origem', '')}"
            if acao == "excluir":
                return None, nota
            if acao == "manual":
                return "__SEM_CATALOGO__", nota
            destino = str(regra.get("destino", "")).strip()
            if destino:
                return destino, nota
            continue

        # Regras avançadas antigas continuam 100% compatíveis.
        todos = [_normalizar(x) for x in regra.get("todos", []) if _normalizar(x)]
        algum = [_normalizar(x) for x in regra.get("algum", []) if _normalizar(x)]
        nenhum = [_normalizar(x) for x in regra.get("nenhum", []) if _normalizar(x)]
        if todos and not all(x in t for x in todos):
            continue
        if algum and not any(x in t for x in algum):
            continue
        if nenhum and any(x in t for x in nenhum):
            continue
        if not todos and not algum:
            continue

        acao = regra.get("acao", "renomear")
        nota = regra.get("nota") or regra.get("nome") or "regra automática"
        if acao == "excluir":
            return None, nota
        if acao == "manual":
            # Migração transparente da regra padrão antiga: o matcher novo entende '/'
            # e não precisa mais mandar Ninho Corda/Trico para lançamento manual.
            if regra.get("nome") == "Ninho Corda/Trico" and not str(regra.get("destino", "")).strip():
                return nome_bruto, "reconhecimento automático de variante com /"
            return "__SEM_CATALOGO__", nota
        destino = str(regra.get("destino", "")).strip()
        if destino:
            return destino, nota
    return nome_bruto, None


def _termos_csv(texto):
    return [x.strip() for x in str(texto or "").replace(";", ",").split(",") if x.strip()]


class ProdutoSelectorDialog(tk.Toplevel):
    """Seletor pesquisável de um produto carregado das planilhas de catálogo."""

    def __init__(self, parent, produtos, selecionado=""):
        super().__init__(parent)
        self.title("Selecionar produto da planilha")
        self.geometry("760x540")
        self.minsize(620, 420)
        self.transient(parent)
        self.grab_set()
        self.resultado = None
        self.produtos = [p for p in (produtos or []) if str(p.get("nome", "")).strip()]
        self.var_busca = tk.StringVar()
        self._after_busca = None

        root = ttk.Frame(self, padding=14)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(2, weight=1)

        ttk.Label(root, text="Escolha o produto exatamente como ele aparece na planilha", style="Section.TLabel").grid(row=0, column=0, sticky="w")
        busca = ttk.Entry(root, textvariable=self.var_busca)
        busca.grid(row=1, column=0, sticky="ew", pady=(8, 8))
        busca.bind("<KeyRelease>", self._agendar_filtro)

        cols = ("produto", "arquivo", "aba")
        self.tree = ttk.Treeview(root, columns=cols, show="headings", selectmode="browse")
        self.tree.heading("produto", text="Produto")
        self.tree.heading("arquivo", text="Planilha")
        self.tree.heading("aba", text="Aba")
        self.tree.column("produto", width=360, anchor="w")
        self.tree.column("arquivo", width=170, anchor="w")
        self.tree.column("aba", width=150, anchor="w")
        self.tree.grid(row=2, column=0, sticky="nsew")
        self.tree.bind("<Double-1>", lambda e: self._confirmar())
        self.tree.bind("<Return>", lambda e: self._confirmar())

        sb = ttk.Scrollbar(root, orient="vertical", command=self.tree.yview)
        sb.grid(row=2, column=1, sticky="ns")
        self.tree.configure(yscrollcommand=sb.set)

        rodape = ttk.Frame(root)
        rodape.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(10, 0))
        self.lbl_qtd = ttk.Label(rodape, text="", style="Muted.TLabel")
        self.lbl_qtd.pack(side="left")
        ttk.Button(rodape, text="Cancelar", command=self.destroy).pack(side="right")
        ttk.Button(rodape, text="Usar este produto", style="Success.TButton", command=self._confirmar).pack(side="right", padx=(0, 8))

        self._render(self.produtos, selecionado)
        busca.focus_set()
        self.bind("<Escape>", lambda e: self.destroy())

    def _agendar_filtro(self, _event=None):
        if self._after_busca:
            try: self.after_cancel(self._after_busca)
            except Exception: pass
        self._after_busca = self.after(90, self._filtrar)

    def _filtrar(self):
        self._after_busca = None
        q = _normalizar(self.var_busca.get())
        if not q:
            filtrados = self.produtos
        else:
            termos = q.split()
            filtrados = [p for p in self.produtos if all(t in _normalizar(p.get("nome", "")) for t in termos)]
        self._render(filtrados)

    def _render(self, itens, selecionado=""):
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        sel_iid = None
        alvo = _normalizar(selecionado)
        for i, p in enumerate(itens):
            iid = str(i)
            self.tree.insert("", "end", iid=iid, values=(p.get("nome", ""), p.get("arquivo", ""), p.get("aba", "")))
            if alvo and _normalizar(p.get("nome", "")) == alvo and sel_iid is None:
                sel_iid = iid
        if sel_iid is not None:
            self.tree.selection_set(sel_iid)
            self.tree.see(sel_iid)
        elif itens:
            self.tree.selection_set("0")
        self.lbl_qtd.config(text=f"{len(itens)} produto(s)")

    def _confirmar(self):
        sel = self.tree.selection()
        if not sel:
            return
        valores = self.tree.item(sel[0], "values")
        if valores:
            self.resultado = str(valores[0])
            self.destroy()


class MapeamentoDialog(tk.Toplevel):
    """Fluxo simples: texto que chega -> produto real escolhido do catálogo."""

    def __init__(self, parent, produtos, regra=None):
        super().__init__(parent)
        self.title("Atribuir produto")
        self.geometry("650x330")
        self.minsize(590, 300)
        self.transient(parent)
        self.grab_set()
        self.resultado = None
        self.produtos = produtos or []
        regra = copy.deepcopy(regra or {})

        self.var_origem = tk.StringVar(value=regra.get("origem", ""))
        self.var_modo = tk.StringVar(value=regra.get("modo_origem", "contem"))
        self.var_destino = tk.StringVar(value=regra.get("destino", ""))
        self.var_ativo = tk.BooleanVar(value=regra.get("ativo", True))

        root = ttk.Frame(self, padding=18)
        root.pack(fill="both", expand=True)
        root.columnconfigure(1, weight=1)

        ttk.Label(root, text="Atribuição de produto", style="Header.TLabel").grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(root, text="Digite como o produto costuma chegar e escolha o produto correto diretamente da planilha.", style="Muted.TLabel", wraplength=570).grid(row=1, column=0, columnspan=3, sticky="w", pady=(2, 16))

        ttk.Label(root, text="Quando chegar").grid(row=2, column=0, sticky="w", pady=6)
        origem = ttk.Entry(root, textvariable=self.var_origem)
        origem.grid(row=2, column=1, sticky="ew", pady=6)
        ttk.Combobox(root, textvariable=self.var_modo, values=["contem", "exato"], state="readonly", width=10).grid(row=2, column=2, padx=(8, 0), pady=6)

        ttk.Label(root, text="Vai virar").grid(row=3, column=0, sticky="w", pady=6)
        ttk.Entry(root, textvariable=self.var_destino, state="readonly").grid(row=3, column=1, sticky="ew", pady=6)
        ttk.Button(root, text="Selecionar da planilha…", command=self._selecionar_destino).grid(row=3, column=2, padx=(8, 0), pady=6)

        ttk.Checkbutton(root, text="Atribuição ativa", variable=self.var_ativo).grid(row=4, column=1, sticky="w", pady=(8, 4))

        botoes = ttk.Frame(root)
        botoes.grid(row=5, column=0, columnspan=3, sticky="e", pady=(14, 0))
        ttk.Button(botoes, text="Cancelar", command=self.destroy).pack(side="left", padx=(0, 8))
        ttk.Button(botoes, text="Salvar atribuição", style="Success.TButton", command=self._salvar).pack(side="left")

        origem.focus_set()
        self.bind("<Escape>", lambda e: self.destroy())

    def _selecionar_destino(self):
        if not self.produtos:
            messagebox.showwarning("Catálogo vazio", "Primeiro carregue uma planilha de produtos em Configurações > Regras automáticas.", parent=self)
            return
        d = ProdutoSelectorDialog(self, self.produtos, self.var_destino.get())
        self.wait_window(d)
        if d.resultado:
            self.var_destino.set(d.resultado)

    def _salvar(self):
        origem = self.var_origem.get().strip()
        destino = self.var_destino.get().strip()
        modo = self.var_modo.get().strip() or "contem"
        if not origem:
            messagebox.showwarning("Produto de entrada", "Digite como o produto chega na mensagem.", parent=self)
            return
        if not destino:
            messagebox.showwarning("Produto de destino", "Selecione o produto correto da planilha.", parent=self)
            return
        self.resultado = {
            "tipo": "mapeamento",
            "nome": f"{origem} → {destino}",
            "acao": "renomear",
            "origem": origem,
            "modo_origem": modo,
            "destino": destino,
            "nota": f"atribuição salva: {origem} -> {destino}",
            "ativo": bool(self.var_ativo.get()),
        }
        self.destroy()


class RegraDialog(tk.Toplevel):
    """Editor das regras avançadas. O destino também é escolhido do catálogo."""

    def __init__(self, parent, regra=None, produtos=None):
        super().__init__(parent)
        self.title("Regra avançada")
        self.geometry("690x540")
        self.minsize(620, 490)
        self.transient(parent)
        self.grab_set()
        self.resultado = None
        self.produtos = produtos if produtos is not None else catalogo_regras()
        regra = copy.deepcopy(regra or {})

        self.vars = {
            "nome": tk.StringVar(value=regra.get("nome", "")),
            "acao": tk.StringVar(value=regra.get("acao", "renomear")),
            "todos": tk.StringVar(value=", ".join(regra.get("todos", []))),
            "algum": tk.StringVar(value=", ".join(regra.get("algum", []))),
            "nenhum": tk.StringVar(value=", ".join(regra.get("nenhum", []))),
            "destino": tk.StringVar(value=regra.get("destino", "")),
            "nota": tk.StringVar(value=regra.get("nota", "")),
            "ativo": tk.BooleanVar(value=regra.get("ativo", True)),
        }

        root = ttk.Frame(self, padding=16)
        root.pack(fill="both", expand=True)
        root.columnconfigure(1, weight=1)
        linha = 0
        ttk.Label(root, text="Regra avançada", style="Section.TLabel").grid(row=linha, column=0, columnspan=2, sticky="w")
        linha += 1
        ttk.Label(root, text="Use isto só quando a atribuição simples não for suficiente. Separe termos por vírgula.", style="Muted.TLabel", wraplength=600).grid(row=linha, column=0, columnspan=2, sticky="w", pady=(2, 12))
        linha += 1

        ttk.Label(root, text="Nome da regra").grid(row=linha, column=0, sticky="w", pady=5)
        ttk.Entry(root, textvariable=self.vars["nome"]).grid(row=linha, column=1, sticky="ew", pady=5); linha += 1

        ttk.Label(root, text="Ação").grid(row=linha, column=0, sticky="w", pady=5)
        acao = ttk.Combobox(root, textvariable=self.vars["acao"], state="readonly", values=["renomear", "excluir", "manual"])
        acao.grid(row=linha, column=1, sticky="ew", pady=5); linha += 1

        for rotulo, key in [("Todos estes termos", "todos"), ("Pelo menos um", "algum"), ("Nenhum destes termos", "nenhum")]:
            ttk.Label(root, text=rotulo).grid(row=linha, column=0, sticky="w", pady=5)
            ttk.Entry(root, textvariable=self.vars[key]).grid(row=linha, column=1, sticky="ew", pady=5); linha += 1

        ttk.Label(root, text="Produto de destino").grid(row=linha, column=0, sticky="w", pady=5)
        dest = ttk.Frame(root)
        dest.grid(row=linha, column=1, sticky="ew", pady=5)
        dest.columnconfigure(0, weight=1)
        ttk.Entry(dest, textvariable=self.vars["destino"], state="readonly").grid(row=0, column=0, sticky="ew")
        ttk.Button(dest, text="Selecionar da planilha…", command=self._selecionar_destino).grid(row=0, column=1, padx=(6, 0))
        linha += 1

        ttk.Label(root, text="Nota exibida").grid(row=linha, column=0, sticky="w", pady=5)
        ttk.Entry(root, textvariable=self.vars["nota"]).grid(row=linha, column=1, sticky="ew", pady=5); linha += 1

        ttk.Checkbutton(root, text="Regra ativa", variable=self.vars["ativo"]).grid(row=linha, column=1, sticky="w", pady=(8, 12)); linha += 1
        botoes = ttk.Frame(root)
        botoes.grid(row=linha, column=0, columnspan=2, sticky="e")
        ttk.Button(botoes, text="Cancelar", command=self.destroy).pack(side="left", padx=(0, 8))
        ttk.Button(botoes, text="Salvar regra", style="Accent.TButton", command=self._salvar).pack(side="left")
        self.bind("<Escape>", lambda e: self.destroy())

    def _selecionar_destino(self):
        if not self.produtos:
            messagebox.showwarning("Catálogo vazio", "Primeiro carregue uma planilha de produtos na aba Regras automáticas.", parent=self)
            return
        d = ProdutoSelectorDialog(self, self.produtos, self.vars["destino"].get())
        self.wait_window(d)
        if d.resultado:
            self.vars["destino"].set(d.resultado)

    def _salvar(self):
        nome = self.vars["nome"].get().strip()
        acao = self.vars["acao"].get().strip()
        destino = self.vars["destino"].get().strip()
        todos = _termos_csv(self.vars["todos"].get())
        algum = _termos_csv(self.vars["algum"].get())
        nenhum = _termos_csv(self.vars["nenhum"].get())
        if not nome:
            messagebox.showwarning("Nome obrigatório", "Dê um nome para a regra.", parent=self); return
        if not todos and not algum:
            messagebox.showwarning("Condição vazia", "Informe ao menos um termo em 'Todos' ou 'Pelo menos um'.", parent=self); return
        if acao == "renomear" and not destino:
            messagebox.showwarning("Destino obrigatório", "Selecione o produto de destino da planilha.", parent=self); return
        self.resultado = {
            "nome": nome, "acao": acao, "todos": todos, "algum": algum, "nenhum": nenhum,
            "destino": destino, "nota": self.vars["nota"].get().strip(), "ativo": bool(self.vars["ativo"].get()),
        }
        self.destroy()


class ConfiguracoesFrame(ttk.Frame):
    """Editor visual das configurações persistentes."""

    def __init__(self, parent):
        super().__init__(parent)
        self._grupo_atual = None
        self._regra_editando_idx = None
        self._carregando_catalogo = False
        self._build()
        self.recarregar()

    def _emitir_alteracao(self):
        self.event_generate("<<ConfiguracaoAlterada>>", when="tail")

    def _build(self):
        cab = ttk.Frame(self, padding=(16, 14, 16, 8))
        cab.pack(fill="x")
        ttk.Label(cab, text="Configurações", style="Header.TLabel").pack(anchor="w")
        ttk.Label(cab, text=f"Salvas automaticamente em {CONFIG_PATH}", style="Muted.TLabel", wraplength=900).pack(anchor="w", pady=(2, 0))

        actions = ttk.Frame(self, padding=(16, 0, 16, 8))
        actions.pack(fill="x")
        ttk.Button(actions, text="Exportar backup", command=self._exportar).pack(side="left")
        ttk.Button(actions, text="Importar backup", command=self._importar).pack(side="left", padx=6)
        ttk.Button(actions, text="Restaurar padrões", style="Danger.TButton", command=self._resetar).pack(side="right")

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self._build_grupos()
        self._build_exclusoes()
        self._build_regras()
        self._build_preferencias()

    def _build_grupos(self):
        tab = ttk.Frame(self.nb, padding=12)
        self.nb.add(tab, text="Grupos de funcionários")
        tab.columnconfigure(1, weight=1)
        tab.rowconfigure(1, weight=1)

        left = ttk.Frame(tab)
        left.grid(row=0, column=0, rowspan=2, sticky="nsew", padx=(0, 12))
        ttk.Label(left, text="Grupos", style="Section.TLabel").pack(anchor="w", pady=(0, 6))
        self.lst_grupos = tk.Listbox(left, width=25, exportselection=False)
        self.lst_grupos.pack(fill="both", expand=True)
        try:
            import tema
            tema.estilizar_listbox(self.lst_grupos)
        except Exception:
            pass
        self.lst_grupos.bind("<<ListboxSelect>>", self._grupo_selecionado)
        bf = ttk.Frame(left)
        bf.pack(fill="x", pady=(8, 0))
        ttk.Button(bf, text="Novo", command=self._novo_grupo).pack(side="left", fill="x", expand=True)
        ttk.Button(bf, text="Excluir", style="Danger.TButton", command=self._excluir_grupo).pack(side="left", fill="x", expand=True, padx=(6, 0))

        right = ttk.Frame(tab)
        right.grid(row=0, column=1, rowspan=2, sticky="nsew")
        right.columnconfigure(0, weight=1)
        right.rowconfigure(3, weight=1)
        ttk.Label(right, text="Nome do grupo").grid(row=0, column=0, sticky="w")
        self.var_grupo_nome = tk.StringVar()
        ttk.Entry(right, textvariable=self.var_grupo_nome).grid(row=1, column=0, sticky="ew", pady=(4, 10))
        ttk.Label(right, text="Funcionários — um por linha").grid(row=2, column=0, sticky="w")
        self.txt_grupo_membros = scrolledtext.ScrolledText(right, wrap="none", height=15)
        self.txt_grupo_membros.grid(row=3, column=0, sticky="nsew", pady=(4, 10))
        try:
            import tema
            tema.estilizar_texto(self.txt_grupo_membros)
        except Exception:
            pass
        ttk.Button(right, text="Salvar alterações do grupo", style="Accent.TButton", command=self._salvar_grupo).grid(row=4, column=0, sticky="e")

    def _build_exclusoes(self):
        tab = ttk.Frame(self.nb, padding=12)
        self.nb.add(tab, text="Exclusões permanentes")
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(1, weight=1)
        ttk.Label(tab, text="Produtos salvos aqui são ignorados automaticamente nas próximas execuções.", style="Muted.TLabel", wraplength=800).grid(row=0, column=0, sticky="w", pady=(0, 8))
        self.tree_exclusoes = ttk.Treeview(tab, columns=("texto", "modo"), show="headings", selectmode="extended")
        self.tree_exclusoes.heading("texto", text="Texto / produto")
        self.tree_exclusoes.heading("modo", text="Correspondência")
        self.tree_exclusoes.column("texto", width=520)
        self.tree_exclusoes.column("modo", width=130, anchor="center")
        self.tree_exclusoes.grid(row=1, column=0, sticky="nsew")

        add = ttk.Frame(tab)
        add.grid(row=2, column=0, sticky="ew", pady=(10, 0))
        add.columnconfigure(0, weight=1)
        self.var_exclusao_texto = tk.StringVar()
        self.var_exclusao_modo = tk.StringVar(value="exato")
        ttk.Entry(add, textvariable=self.var_exclusao_texto).grid(row=0, column=0, sticky="ew")
        ttk.Combobox(add, textvariable=self.var_exclusao_modo, values=["exato", "contem"], state="readonly", width=12).grid(row=0, column=1, padx=6)
        ttk.Button(add, text="Adicionar", style="Accent.TButton", command=self._adicionar_exclusao).grid(row=0, column=2)
        ttk.Button(add, text="Remover selecionados", style="Danger.TButton", command=self._remover_exclusoes).grid(row=0, column=3, padx=(6, 0))

    def _build_regras(self):
        tab = ttk.Frame(self.nb, padding=12)
        self.nb.add(tab, text="Regras automáticas")
        tab.columnconfigure(0, weight=1)
        tab.rowconfigure(3, weight=1)

        catalogo_box = ttk.LabelFrame(tab, text=" Catálogo de produtos ", padding=10)
        catalogo_box.grid(row=0, column=0, sticky="ew")
        catalogo_box.columnconfigure(0, weight=1)
        self.var_catalogo_status = tk.StringVar(value="Nenhuma planilha configurada")
        ttk.Label(catalogo_box, textvariable=self.var_catalogo_status, style="Muted.TLabel", wraplength=700).grid(row=0, column=0, sticky="w")
        cb = ttk.Frame(catalogo_box)
        cb.grid(row=0, column=1, sticky="e", padx=(10, 0))
        self.btn_catalogo_escolher = ttk.Button(cb, text="Escolher planilha(s)…", command=self._escolher_planilhas_catalogo)
        self.btn_catalogo_escolher.pack(side="left")
        self.btn_catalogo_atualizar = ttk.Button(cb, text="Atualizar", command=self._atualizar_catalogo_salvo)
        self.btn_catalogo_atualizar.pack(side="left", padx=(6, 0))
        ttk.Label(catalogo_box, text="O caminho e os produtos ficam salvos. Pode selecionar Cadeiras e Mesas juntas.", style="Muted.TLabel").grid(row=1, column=0, columnspan=2, sticky="w", pady=(4, 0))

        topo = ttk.Frame(tab)
        topo.grid(row=1, column=0, sticky="ew", pady=(10, 6))
        ttk.Label(topo, text="Atribuições salvas", style="Section.TLabel").pack(side="left")
        ttk.Button(topo, text="+ Atribuir produto", style="Success.TButton", command=self._novo_mapeamento).pack(side="right")
        ttk.Button(topo, text="+ Regra avançada", command=self._nova_regra_avancada).pack(side="right", padx=(0, 6))

        ttk.Label(tab, text="Ex.: ‘Mesa Moscou’ → escolha ‘MESA HAPPY HOUR’ da planilha. As regras de cima têm prioridade.", style="Muted.TLabel").grid(row=2, column=0, sticky="w", pady=(0, 6))

        self.tree_regras = ttk.Treeview(tab, columns=("ativo", "nome", "acao", "cond", "destino"), show="headings", selectmode="browse")
        for c, t, w in [("ativo", "Ativa", 55), ("nome", "Regra", 180), ("acao", "Ação", 90), ("cond", "Quando chegar", 310), ("destino", "Vira / Destino", 230)]:
            self.tree_regras.heading(c, text=t)
            self.tree_regras.column(c, width=w, anchor="w")
        self.tree_regras.grid(row=3, column=0, sticky="nsew")
        self.tree_regras.bind("<Double-1>", lambda e: self._editar_regra())
        sb = ttk.Scrollbar(tab, orient="vertical", command=self.tree_regras.yview)
        sb.grid(row=3, column=1, sticky="ns")
        self.tree_regras.configure(yscrollcommand=sb.set)

        b = ttk.Frame(tab)
        b.grid(row=4, column=0, sticky="ew", pady=(10, 0))
        ttk.Button(b, text="Editar", command=self._editar_regra).pack(side="left")
        ttk.Button(b, text="Excluir", style="Danger.TButton", command=self._excluir_regra).pack(side="left", padx=6)
        ttk.Button(b, text="↑", width=3, command=lambda: self._mover_regra(-1)).pack(side="right", padx=(6, 0))
        ttk.Button(b, text="↓", width=3, command=lambda: self._mover_regra(1)).pack(side="right")

    def _build_preferencias(self):
        tab = ttk.Frame(self.nb, padding=16)
        self.nb.add(tab, text="Preferências")
        tab.columnconfigure(1, weight=1)
        self.var_aceitar = tk.IntVar(value=85)
        self.var_revisar = tk.IntVar(value=60)
        self.var_funcionario = tk.IntVar(value=65)
        self.var_somar = tk.BooleanVar(value=True)
        self.var_lembrar = tk.BooleanVar(value=True)
        self.var_confirmar = tk.BooleanVar(value=True)

        ttk.Label(tab, text="Reconhecimento automático", style="Section.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        campos = [("Aceitar como confiável", self.var_aceitar), ("Enviar para revisão abaixo de", self.var_revisar), ("Reconhecer funcionário a partir de", self.var_funcionario)]
        for i, (rot, var) in enumerate(campos, start=1):
            ttk.Label(tab, text=rot).grid(row=i, column=0, sticky="w", pady=5)
            ttk.Spinbox(tab, from_=0, to=100, textvariable=var, width=7).grid(row=i, column=1, sticky="w", pady=5)
        r = len(campos) + 2
        ttk.Separator(tab).grid(row=r, column=0, columnspan=2, sticky="ew", pady=12)
        r += 1
        ttk.Checkbutton(tab, text="Somar com valores já existentes por padrão", variable=self.var_somar).grid(row=r, column=0, columnspan=2, sticky="w", pady=4); r += 1
        ttk.Checkbutton(tab, text="Lembrar últimos arquivos e pasta de saída", variable=self.var_lembrar).grid(row=r, column=0, columnspan=2, sticky="w", pady=4); r += 1
        ttk.Checkbutton(tab, text="Pedir confirmação antes de gravar planilhas", variable=self.var_confirmar).grid(row=r, column=0, columnspan=2, sticky="w", pady=4); r += 1
        ttk.Button(tab, text="Salvar preferências", style="Accent.TButton", command=self._salvar_preferencias).grid(row=r, column=0, columnspan=2, sticky="w", pady=(14, 0))

    def recarregar(self):
        data = STORE.data()
        self._refresh_grupos(data.get("grupos", {}))
        self._refresh_exclusoes(data.get("exclusoes_produtos", []))
        self._refresh_regras(data.get("regras_produtos", []))
        self._atualizar_status_catalogo(data.get("catalogo_regras", {}))
        match = data.get("matching", {})
        self.var_aceitar.set(int(match.get("aceitar", 85)))
        self.var_revisar.set(int(match.get("revisar", 60)))
        self.var_funcionario.set(int(match.get("funcionario", 65)))
        pref = data.get("preferencias", {})
        self.var_somar.set(bool(pref.get("somar_existente", True)))
        self.var_lembrar.set(bool(pref.get("lembrar_arquivos", True)))
        self.var_confirmar.set(bool(pref.get("confirmar_salvamento", True)))

    def _refresh_grupos(self, grupos):
        atual = self._grupo_atual
        self.lst_grupos.delete(0, tk.END)
        for nome in grupos:
            self.lst_grupos.insert(tk.END, nome)
        if atual in grupos:
            idx = list(grupos).index(atual)
            self.lst_grupos.selection_set(idx)
            self.lst_grupos.see(idx)
            self._mostrar_grupo(atual)
        elif grupos:
            self.lst_grupos.selection_set(0)
            self._mostrar_grupo(next(iter(grupos)))
        else:
            self._grupo_atual = None
            self.var_grupo_nome.set("")
            self.txt_grupo_membros.delete("1.0", tk.END)

    def _grupo_selecionado(self, _e=None):
        sel = self.lst_grupos.curselection()
        if sel:
            self._mostrar_grupo(self.lst_grupos.get(sel[0]))

    def _mostrar_grupo(self, nome):
        self._grupo_atual = nome
        grupos = STORE.get("grupos", default={}) or {}
        self.var_grupo_nome.set(nome)
        self.txt_grupo_membros.delete("1.0", tk.END)
        self.txt_grupo_membros.insert("1.0", "\n".join(grupos.get(nome, [])))

    def _novo_grupo(self):
        grupos = STORE.get("grupos", default={}) or {}
        base = "Novo grupo"
        nome = base
        n = 2
        while nome in grupos:
            nome = f"{base} {n}"; n += 1
        grupos[nome] = []
        STORE.replace_section("grupos", grupos)
        self._grupo_atual = nome
        self._refresh_grupos(grupos)
        self._emitir_alteracao()

    def _salvar_grupo(self):
        if self._grupo_atual is None:
            return
        novo_nome = self.var_grupo_nome.get().strip()
        if not novo_nome:
            messagebox.showwarning("Nome obrigatório", "Informe o nome do grupo.")
            return
        membros = []
        vistos = set()
        for linha in self.txt_grupo_membros.get("1.0", "end").splitlines():
            nome = linha.strip()
            if nome and _normalizar(nome) not in vistos:
                membros.append(nome)
                vistos.add(_normalizar(nome))
        grupos = STORE.get("grupos", default={}) or {}
        if novo_nome != self._grupo_atual and novo_nome in grupos:
            messagebox.showwarning("Grupo já existe", "Já existe um grupo com esse nome.")
            return
        novo = {}
        for nome, vals in grupos.items():
            if nome == self._grupo_atual:
                novo[novo_nome] = membros
            else:
                novo[nome] = vals
        if self._grupo_atual not in grupos:
            novo[novo_nome] = membros
        self._grupo_atual = novo_nome
        STORE.replace_section("grupos", novo)
        self._refresh_grupos(novo)
        self._emitir_alteracao()

    def _excluir_grupo(self):
        if not self._grupo_atual:
            return
        if not messagebox.askyesno("Excluir grupo", f"Excluir o grupo '{self._grupo_atual}'?"):
            return
        grupos = STORE.get("grupos", default={}) or {}
        grupos.pop(self._grupo_atual, None)
        self._grupo_atual = None
        STORE.replace_section("grupos", grupos)
        self._refresh_grupos(grupos)
        self._emitir_alteracao()

    def _refresh_exclusoes(self, itens):
        for i in self.tree_exclusoes.get_children():
            self.tree_exclusoes.delete(i)
        for idx, item in enumerate(itens):
            if isinstance(item, str):
                texto, modo = item, "contem"
            else:
                texto, modo = item.get("texto", ""), item.get("modo", "contem")
            self.tree_exclusoes.insert("", "end", iid=str(idx), values=(texto, "Nome exato" if modo == "exato" else "Contém"))

    def _adicionar_exclusao(self):
        texto = self.var_exclusao_texto.get().strip()
        if not texto:
            return
        modo = self.var_exclusao_modo.get() or "exato"
        adicionar_exclusao(texto, modo)
        self.var_exclusao_texto.set("")
        self._refresh_exclusoes(STORE.get("exclusoes_produtos", default=[]))
        self._emitir_alteracao()

    def _remover_exclusoes(self):
        sel = sorted((int(i) for i in self.tree_exclusoes.selection()), reverse=True)
        if not sel:
            return
        itens = STORE.get("exclusoes_produtos", default=[]) or []
        for idx in sel:
            if 0 <= idx < len(itens):
                itens.pop(idx)
        STORE.replace_section("exclusoes_produtos", itens)
        self._refresh_exclusoes(itens)
        self._emitir_alteracao()

    def _atualizar_status_catalogo(self, sec=None):
        sec = sec if isinstance(sec, dict) else (STORE.get("catalogo_regras", default={}) or {})
        paths = sec.get("planilhas", []) or []
        produtos = sec.get("produtos", []) or []
        if not paths:
            texto = "Nenhuma planilha configurada — escolha a planilha que contém os produtos."
        else:
            nomes = ", ".join(os.path.basename(p) for p in paths)
            texto = f"{nomes}  •  {len(produtos)} produto(s) disponíveis e salvos"
        self.var_catalogo_status.set(texto)

    def _escolher_planilhas_catalogo(self):
        paths = filedialog.askopenfilenames(
            title="Selecione a(s) planilha(s) que contêm os produtos",
            filetypes=[("Planilhas Excel", "*.xlsx *.xlsm"), ("Todos os arquivos", "*.*")],
        )
        if not paths:
            return
        self._carregar_catalogo_async(list(paths))

    def _atualizar_catalogo_salvo(self):
        sec = STORE.get("catalogo_regras", default={}) or {}
        paths = sec.get("planilhas", []) or []
        if not paths:
            self._escolher_planilhas_catalogo()
            return
        self._carregar_catalogo_async(paths)

    def _carregar_catalogo_async(self, paths):
        if self._carregando_catalogo:
            return
        self._carregando_catalogo = True
        self.var_catalogo_status.set("Lendo produtos da(s) planilha(s)…")
        self.btn_catalogo_escolher.config(state="disabled")
        self.btn_catalogo_atualizar.config(state="disabled")

        resultados_q = queue.Queue(maxsize=1)

        def worker():
            produtos = []
            erros = []
            vistos = set()
            for path in paths:
                try:
                    for item in extrair_catalogo_planilha(path):
                        chave = _normalizar(item.get("nome", ""))
                        if chave and chave not in vistos:
                            vistos.add(chave)
                            produtos.append(item)
                except Exception as e:
                    erros.append(f"{os.path.basename(path)}: {e}")
            produtos.sort(key=lambda x: _normalizar(x.get("nome", "")))
            if produtos:
                salvar_catalogo_regras(paths, produtos)
            resultados_q.put((produtos, erros))

        def aguardar_worker():
            try:
                produtos, erros = resultados_q.get_nowait()
            except queue.Empty:
                self.after(80, aguardar_worker)
                return
            finalizado(produtos, erros)

        def finalizado(produtos, erros):
            self._carregando_catalogo = False
            self.btn_catalogo_escolher.config(state="normal")
            self.btn_catalogo_atualizar.config(state="normal")
            self._atualizar_status_catalogo()
            if not produtos:
                messagebox.showerror("Nenhum produto encontrado", "Não consegui encontrar produtos nas planilhas selecionadas." + ("\n\n" + "\n".join(erros) if erros else ""))
            elif erros:
                messagebox.showwarning("Catálogo carregado parcialmente", "Os produtos encontrados foram salvos, mas houve problemas:\n\n" + "\n".join(erros))
            else:
                messagebox.showinfo("Catálogo salvo", f"{len(produtos)} produto(s) carregados. Agora você pode selecionar o destino das regras sem digitar o nome.")

        threading.Thread(target=worker, daemon=True).start()
        self.after(80, aguardar_worker)

    def _novo_mapeamento(self):
        produtos = catalogo_regras()
        if not produtos:
            messagebox.showwarning("Catálogo vazio", "Primeiro clique em 'Escolher planilha(s)…' e carregue a planilha que contém os produtos.")
            return
        d = MapeamentoDialog(self, produtos)
        self.wait_window(d)
        if not d.resultado:
            return
        regra = d.resultado
        regras = STORE.get("regras_produtos", default=[]) or []
        duplicado = next((i for i, r in enumerate(regras) if r.get("tipo") == "mapeamento" and _normalizar(r.get("origem", "")) == _normalizar(regra.get("origem", "")) and r.get("modo_origem", "contem") == regra.get("modo_origem", "contem")), None)
        if duplicado is not None:
            if not messagebox.askyesno("Atribuição já existe", "Já existe uma atribuição para esse texto. Deseja substituir pelo novo produto?"):
                return
            regras[duplicado] = regra
            selecionar = duplicado
        else:
            # Atribuições explícitas do usuário ganham prioridade sobre regras genéricas.
            regras.insert(0, regra)
            selecionar = 0
        STORE.replace_section("regras_produtos", regras)
        self._refresh_regras(regras)
        self.tree_regras.selection_set(str(selecionar))
        self.tree_regras.see(str(selecionar))
        self._emitir_alteracao()

    def _cond_regra(self, r):
        if r.get("tipo") == "mapeamento":
            modo = "nome exato" if r.get("modo_origem") == "exato" else "contém"
            return f"{modo}: {r.get('origem', '')}"
        partes = []
        if r.get("todos"): partes.append("todos: " + ", ".join(r["todos"]))
        if r.get("algum"): partes.append("algum: " + ", ".join(r["algum"]))
        if r.get("nenhum"): partes.append("sem: " + ", ".join(r["nenhum"]))
        return " | ".join(partes)

    def _refresh_regras(self, regras):
        for i in self.tree_regras.get_children(): self.tree_regras.delete(i)
        for idx, r in enumerate(regras):
            acao = "atribuir" if r.get("tipo") == "mapeamento" else r.get("acao", "")
            self.tree_regras.insert("", "end", iid=str(idx), values=("Sim" if r.get("ativo", True) else "Não", r.get("nome", ""), acao, self._cond_regra(r), r.get("destino", "")))

    def _nova_regra(self):
        # Compatibilidade com chamadas antigas: abre uma regra avançada.
        self._nova_regra_avancada()

    def _nova_regra_avancada(self):
        d = RegraDialog(self, produtos=catalogo_regras())
        self.wait_window(d)
        if d.resultado:
            regras = STORE.get("regras_produtos", default=[]) or []
            regras.append(d.resultado)
            STORE.replace_section("regras_produtos", regras)
            self._refresh_regras(regras)
            self._emitir_alteracao()

    def _indice_regra(self):
        sel = self.tree_regras.selection()
        return int(sel[0]) if sel else None

    def _editar_regra(self):
        idx = self._indice_regra()
        if idx is None: return
        regras = STORE.get("regras_produtos", default=[]) or []
        if not (0 <= idx < len(regras)): return
        regra = regras[idx]
        if regra.get("tipo") == "mapeamento":
            d = MapeamentoDialog(self, catalogo_regras(), regra)
            self.wait_window(d)
            if d.resultado:
                regras[idx] = d.resultado
                STORE.replace_section("regras_produtos", regras)
                self._refresh_regras(regras)
                self.tree_regras.selection_set(str(idx))
                self._emitir_alteracao()
            return
        d = RegraDialog(self, regra, produtos=catalogo_regras())
        self.wait_window(d)
        if d.resultado:
            regras[idx] = d.resultado
            STORE.replace_section("regras_produtos", regras)
            self._refresh_regras(regras)
            self.tree_regras.selection_set(str(idx))
            self._emitir_alteracao()

    def _excluir_regra(self):
        idx = self._indice_regra()
        if idx is None: return
        regras = STORE.get("regras_produtos", default=[]) or []
        if 0 <= idx < len(regras):
            regras.pop(idx)
            STORE.replace_section("regras_produtos", regras)
            self._refresh_regras(regras)
            self._emitir_alteracao()

    def _mover_regra(self, delta):
        idx = self._indice_regra()
        if idx is None: return
        regras = STORE.get("regras_produtos", default=[]) or []
        novo = idx + delta
        if 0 <= novo < len(regras):
            regras[idx], regras[novo] = regras[novo], regras[idx]
            STORE.replace_section("regras_produtos", regras)
            self._refresh_regras(regras)
            self.tree_regras.selection_set(str(novo))
            self.tree_regras.see(str(novo))
            self._emitir_alteracao()

    def _salvar_preferencias(self):
        aceitar = max(0, min(100, int(self.var_aceitar.get())))
        revisar = max(0, min(100, int(self.var_revisar.get())))
        funcionario = max(0, min(100, int(self.var_funcionario.get())))
        if revisar > aceitar:
            messagebox.showwarning("Limites inválidos", "O limite de revisão não pode ser maior que o limite de aceitação.")
            return
        STORE.replace_section("matching", {"aceitar": aceitar, "revisar": revisar, "funcionario": funcionario})
        STORE.replace_section("preferencias", {"somar_existente": bool(self.var_somar.get()), "lembrar_arquivos": bool(self.var_lembrar.get()), "confirmar_salvamento": bool(self.var_confirmar.get())})
        self._emitir_alteracao()
        messagebox.showinfo("Configurações", "Preferências salvas.")

    def _exportar(self):
        path = filedialog.asksaveasfilename(title="Exportar configuração", defaultextension=".json", filetypes=[("JSON", "*.json")], initialfile="suite_producao_config.json")
        if not path: return
        try:
            STORE.export_to(path)
            messagebox.showinfo("Backup criado", f"Configuração exportada para:\n{path}")
        except Exception as e:
            messagebox.showerror("Erro", str(e))

    def _importar(self):
        path = filedialog.askopenfilename(title="Importar configuração", filetypes=[("JSON", "*.json")])
        if not path: return
        if not messagebox.askyesno("Importar configuração", "Isso substituirá grupos, exclusões, regras e preferências atuais. Continuar?"):
            return
        try:
            STORE.import_from(path)
            self.recarregar()
            self._emitir_alteracao()
        except Exception as e:
            messagebox.showerror("Configuração inválida", str(e))

    def _resetar(self):
        if not messagebox.askyesno("Restaurar padrões", "Restaurar todas as configurações originais? Suas personalizações serão apagadas."):
            return
        STORE.reset()
        self._grupo_atual = None
        self._regra_editando_idx = None
        self.recarregar()
        self._emitir_alteracao()

