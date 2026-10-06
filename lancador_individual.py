"""Lançador mensal por funcionário — otimizado, editável e não bloqueante."""

from __future__ import annotations

import os
import re
import unicodedata
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

import openpyxl
from rapidfuzz import fuzz

try:
    from . import tema, configuracao
except ImportError:
    import tema
    import configuracao

PRODUTO_COL = 1
DIA_HEADER_ROW = 2
SHEETS_IGNORADAS = ("TOTAL GERAL",)

# Valores padrão mantidos como API; a execução usa os limites editáveis da configuração.
ACCEPT_THRESHOLD = 85
REVIEW_THRESHOLD = 60
FUNCIONARIO_THRESHOLD = 65


def _thresholds():
    return configuracao.thresholds()


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = texto.upper()
    texto = re.sub(r"[^A-Z0-9 ]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def pontuar_texto(consulta_tokens, alvo_norm, alvo_tokens):
    if not consulta_tokens or not alvo_tokens:
        return 0.0
    tscores = [max(fuzz.ratio(qt, tt) for tt in alvo_tokens) for qt in consulta_tokens]
    return 0.7 * (sum(tscores) / len(tscores)) + 0.3 * fuzz.token_set_ratio(" ".join(consulta_tokens), alvo_norm)


def melhor_match_lista(nome_bruto, itens_norm_tokens):
    consulta = normalizar(nome_bruto)
    tokens = consulta.split()
    if not tokens or not itens_norm_tokens:
        return None, 0.0
    melhor, melhor_score = None, -1.0
    for chave, norm, alvo_tokens in itens_norm_tokens:
        score = pontuar_texto(tokens, norm, alvo_tokens)
        if score > melhor_score:
            melhor, melhor_score = chave, score
    return melhor, melhor_score


# Matching robusto usado SOMENTE para produtos. Funcionários mantêm a lógica própria.
# Ajustado e testado contra as planilhas reais de agosto/2026 enviadas pelo usuário.

_TIPOS_PRODUTO = {
    "CADE": "CADE", "CADEIRA": "CADE", "CADEIRAS": "CADE",
    "MESA": "MESA", "MESAS": "MESA",
    "POLT": "POLT", "POLTRONA": "POLT", "POLTRONAS": "POLT",
    "BANQ": "BANQ", "BANQUETA": "BANQ", "BANQUETAS": "BANQ",
    "NAMO": "NAMO", "NAMORADEIRA": "NAMO", "NAMORADEIRAS": "NAMO",
    "BALANCO": "BALANCO", "BALANCOS": "BALANCO",
    "CHAISE": "CHAISE", "CHAISES": "CHAISE",
    "ESPR": "ESPR", "ESP": "ESPR", "ESPREGUICADEIRA": "ESPR", "ESPREGUICADEIRAS": "ESPR",
    "BISTRO": "BISTRO",
    "CACH": "CACH", "CACHEPO": "CACH", "CACHEPOS": "CACH",
    "PUFF": "PUFF",
    "ASSENTO": "ASSENTO", "ENCOSTO": "ENCOSTO",
    "INFANTIL": "INFANTIL", "DESCANSO": "DESCANSO",
}

_SINONIMOS_PRODUTO = {
    "CONF": "CONFORTO", "CONFORT": "CONFORTO",
    "BERLIN": "BERLIM",
    "NAVE": "NAVEGANTES", "NAVEGANTE": "NAVEGANTES", "NAVEGANTES": "NAVEGANTES",
    "TRAM": "TRAMADO", "TRAMADA": "TRAMADO",
    "ALU": "ALUMINIO", "ALUM": "ALUMINIO",
    "NAUTICA": "NAUTICO",
    "CORDAO": "CORDA",
    "OBI": "OBILONGO",
    "INF": "INFANTIL",
}

# Acabamento/material ajuda a escolher a variante, mas nunca pode valer mais que o modelo.
_DETALHES_PRODUTO = {
    "TRICO", "CORDA", "FIBRA", "ARG", "ARGOLA", "TRAMADO", "ALUMINIO", "FERRO",
    "CRA", "ENC", "ARTESANAL", "BAIXA", "BAIXO", "FIO", "COLMEIA", "NAUTICO",
    "PINTADA", "POLIDA", "FLECHA", "OBILONGO", "RET", "QUAD", "COM", "SEM",
    "BRACO", "BRACOS", "ENCAPAGEM", "BASE", "PALITO", "SO", "COMPLETA", "COMPLETO",
}
_STOP_PRODUTO = {"DE", "DO", "DA", "DOS", "DAS", "A", "O", "E", "NA", "NO", "EM"}


def _normalizar_produto_barra(texto):
    """Normaliza preservando '/', porque a barra tem significado nos nomes reais."""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().upper()
    texto = texto.replace("\\", "/")
    texto = re.sub(r"\b([468])\s*LUG(?:ARES)?\b", r"\1LUG", texto)
    texto = re.sub(r"\b([468])\s*L\b", r"\1LUG", texto)
    texto = re.sub(r"\s*/\s*", "/", texto)
    texto = re.sub(r"[^A-Z0-9/ ]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _canon_produto(token):
    return _SINONIMOS_PRODUTO.get(token, token)


def _variantes_produto(texto):
    """Gera variantes de '/' sem confundir separador com alias.

    TAUPE/ELEGANTE e PANERO/VIENA geram aliases. Já ASSENTO/HARMONIA,
    ENCOSTO/IPANEMA e FIBRA/SEM não perdem a primeira parte, porque ali a barra
    funciona como separador estrutural, não como dois nomes equivalentes.
    """
    bruto = _normalizar_produto_barra(texto)
    if not bruto:
        return []
    grupos, achatada = [], []
    for parte in bruto.split():
        alts = [_canon_produto(x) for x in parte.split("/") if x]
        if not alts:
            continue
        achatada.extend(alts)
        aliasavel = len(alts) > 1 and all(
            x not in _TIPOS_PRODUTO and x not in _DETALHES_PRODUTO and x not in _STOP_PRODUTO
            and not re.fullmatch(r"[468]LUG", x)
            for x in alts
        )
        grupos.append([[x] for x in alts] if aliasavel else [alts])

    variantes = [achatada] if achatada else []
    combos = [[]]
    for escolhas in grupos:
        combos = [prefixo + escolha for prefixo in combos for escolha in escolhas][:64]
    variantes.extend(combos)

    unicas, vistos = [], set()
    for variante in variantes:
        chave = tuple(variante)
        if chave and chave not in vistos:
            vistos.add(chave)
            unicas.append(variante)
    return unicas


def _partes_produto(tokens):
    """Retorna (tipo/família, identidade/modelo, detalhes)."""
    tipo = None
    resto = []
    for i, tok in enumerate(list(tokens or [])):
        if tok == "BASE":
            continue
        if tipo is None and tok in _TIPOS_PRODUTO and (i <= 1 or not resto):
            tipo = _TIPOS_PRODUTO[tok]
            continue
        resto.append(tok)

    identidade, detalhes = [], []
    for tok in resto:
        if tok in _STOP_PRODUTO:
            continue
        if tok in _DETALHES_PRODUTO or re.fullmatch(r"[468]LUG", tok):
            detalhes.append(tok)
        elif tok in _TIPOS_PRODUTO and tipo is None:
            tipo = _TIPOS_PRODUTO[tok]
        else:
            identidade.append(tok)
    return tipo, identidade, detalhes


def _media_fuzzy(origem, alvo):
    if not origem or not alvo:
        return 0.0
    return sum(max(fuzz.ratio(x, y) for y in alvo) for x in origem) / len(origem)


def _similaridade_identidade(q, a):
    if not q or not a:
        return 0.0
    sq, sa = set(q), set(a)
    if sq == sa:
        return 100.0
    qtxt, atxt = " ".join(q), " ".join(a)
    direcao_q = _media_fuzzy(q, a)
    direcao_a = _media_fuzzy(a, q)
    score = (
        0.42 * fuzz.token_set_ratio(qtxt, atxt)
        + 0.28 * direcao_q
        + 0.12 * direcao_a
        + 0.18 * fuzz.ratio(qtxt, atxt)
    )
    # "GOTA" precisa casar fortemente com "NINHO GOTA", por exemplo.
    if sq.issubset(sa) or sa.issubset(sq):
        score = max(score, 96.0)
    elif sq & sa:
        score = max(score, 88.0)
    elif max(direcao_q, direcao_a) < 78:
        score = min(score, 58.0)
    return min(100.0, score)


def _similaridade_detalhes(q, a):
    if not q:
        return 100.0
    if not a:
        return 25.0
    sq, sa = set(q), set(a)
    if sq == sa:
        return 100.0
    if sq.issubset(sa):
        return 98.0
    overlap = len(sq & sa) / len(sq)
    return (
        0.50 * fuzz.token_set_ratio(" ".join(q), " ".join(a))
        + 0.25 * (100.0 * overlap)
        + 0.25 * _media_fuzzy(q, a)
    )


def _pontuar_variante_produto(consulta_tokens, alvo_tokens):
    if not consulta_tokens or not alvo_tokens:
        return 0.0
    qtxt, atxt = " ".join(consulta_tokens), " ".join(alvo_tokens)
    if qtxt == atxt:
        return 100.0

    tipo_q, ident_q, detalhe_q = _partes_produto(consulta_tokens)
    tipo_a, ident_a, detalhe_a = _partes_produto(alvo_tokens)
    if not ident_q or not ident_a:
        return 0.65 * fuzz.token_set_ratio(qtxt, atxt) + 0.35 * fuzz.ratio(qtxt, atxt)

    ident = _similaridade_identidade(ident_q, ident_a)
    detalhes = _similaridade_detalhes(detalhe_q, detalhe_a)
    completo = 0.65 * fuzz.token_set_ratio(qtxt, atxt) + 0.35 * fuzz.ratio(qtxt, atxt)

    # Tipo explícito incompatível praticamente elimina o candidato.
    if tipo_q and tipo_a and tipo_q != tipo_a:
        return min(38.0, 0.75 * ident + 0.15 * detalhes + 0.10 * completo)

    score = 0.78 * ident + 0.14 * detalhes + 0.08 * completo
    if tipo_q and tipo_a and tipo_q == tipo_a:
        score += 2.0

    sq, sa = set(ident_q), set(ident_a)
    if sq == sa and (not detalhe_q or set(detalhe_q).issubset(set(detalhe_a))):
        score = max(score, 97.0)
    elif (sq.issubset(sa) or sa.issubset(sq)) and (not detalhe_q or detalhes >= 80):
        score = max(score, 92.0)

    # Uma palavra de acabamento igual não pode salvar um modelo diferente.
    if not (sq & sa) and ident < 82:
        score = min(score, 55.0)
    return max(0.0, min(100.0, score))


def pontuar_nome_produto(nome_consulta, nome_alvo):
    """Pontuação hierárquica e tolerante a acento, abreviação, typo e '/'."""
    qnorm = _normalizar_produto_barra(nome_consulta).replace("/", " ")
    anorm = _normalizar_produto_barra(nome_alvo).replace("/", " ")
    if qnorm and qnorm == anorm:
        return 100.0
    qvars, avars = _variantes_produto(nome_consulta), _variantes_produto(nome_alvo)
    if not qvars or not avars:
        return 0.0
    return max(_pontuar_variante_produto(q, a) for q in qvars for a in avars)


def _melhor_match_produto_avaliado(nome_bruto, pares):
    """Retorna (chave, score) e derruba empate perigoso para pendência."""
    if not _normalizar_produto_barra(nome_bruto) or not pares:
        return None, 0.0

    qnorm = _normalizar_produto_barra(nome_bruto).replace("/", " ")
    for chave, nome_real in pares:
        if _normalizar_produto_barra(nome_real).replace("/", " ") == qnorm:
            return chave, 100.0

    avaliados = sorted(
        ((pontuar_nome_produto(nome_bruto, nome_real), chave, nome_real) for chave, nome_real in pares),
        key=lambda x: x[0], reverse=True,
    )
    melhor_score, melhor, _ = avaliados[0]
    if len(avaliados) > 1:
        segundo_score = avaliados[1][0]
        # Empate/ambiguidade real: não lança no chute; fica pendente (< limite de revisão).
        if segundo_score >= 72 and (melhor_score - segundo_score) < 4.0:
            melhor_score = min(melhor_score, 59.0)
    return melhor, melhor_score


def pontuar_produto(consulta_tokens, alvo_norm, alvo_tokens):
    return pontuar_nome_produto(" ".join(consulta_tokens or []), alvo_norm)


def melhor_match_produto_lista(nome_bruto, itens_norm_tokens):
    pares = []
    for item in itens_norm_tokens:
        if len(item) >= 4:
            chave, nome_real = item[0], item[1]
        else:
            chave, nome_real = item[0], item[1]
        pares.append((chave, nome_real))
    return _melhor_match_produto_avaliado(nome_bruto, pares)


def sheet_tem_dias(ws):
    # Compatível com o layout antigo (B1 contém DIA), com fallback tolerante.
    for r in range(1, min(3, ws.max_row or 1) + 1):
        for c in range(1, min(5, ws.max_column or 1) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "DIA" in v.upper().replace(" ", ""):
                return True
    return False


def carregar_dia_colunas(ws):
    dia_col = {}
    for c in range(2, (ws.max_column or 1) + 1):
        v = ws.cell(row=DIA_HEADER_ROW, column=c).value
        if isinstance(v, int) and 1 <= v <= 31:
            dia_col[v] = c
        elif isinstance(v, float) and v.is_integer() and 1 <= int(v) <= 31:
            dia_col[int(v)] = c
    return dia_col


def carregar_catalogo_funcionario(ws):
    catalogo = []
    for r in range(DIA_HEADER_ROW + 1, (ws.max_row or DIA_HEADER_ROW) + 1):
        nome = ws.cell(row=r, column=PRODUTO_COL).value
        if not isinstance(nome, str) or not nome.strip():
            continue
        nome = nome.strip()
        if nome.upper() in ("TOTAIS", "TOTAL"):
            continue
        norm = normalizar(nome)
        tokens = norm.split()
        if tokens:
            catalogo.append({"row": r, "name": nome, "norm": norm, "tokens": tokens})
    return catalogo


def carregar_estrutura_mensal(caminho, progresso=None):
    """Abre a planilha UMA vez e carrega nomes, catálogos e colunas de dias."""
    wb = openpyxl.load_workbook(caminho, read_only=False, data_only=False,
                                keep_vba=str(caminho).lower().endswith(".xlsm"))
    try:
        nomes = []
        for sn in wb.sheetnames:
            if any(sn.upper().startswith(p) for p in SHEETS_IGNORADAS):
                continue
            if sheet_tem_dias(wb[sn]):
                nomes.append(sn)
        catalogos, dia_cols = {}, {}
        total = max(1, len(nomes))
        for i, sn in enumerate(nomes, start=1):
            ws = wb[sn]
            catalogos[sn] = carregar_catalogo_funcionario(ws)
            dia_cols[sn] = carregar_dia_colunas(ws)
            if progresso:
                progresso(i, total)
        return nomes, catalogos, dia_cols
    finally:
        wb.close()


def listar_funcionarios(caminho):
    nomes, _, _ = carregar_estrutura_mensal(caminho)
    return nomes


RE_DATA = re.compile(r"^(\d{1,2})/(\d{1,2})$")
RE_PRODUTO = re.compile(r"^(\d+)\s+(.+?)\s*$")
RE_PREFIXO_WHATSAPP = re.compile(r"^\[\d{1,2}:\d{2}(?::\d{2})?,\s*[^\]]+\]\s*[^:]*:\s*")


def parse_texto(texto, nomes_funcionarios):
    candidatos = [(n, normalizar(n), normalizar(n).split()) for n in nomes_funcionarios]
    _, _, funcionario_threshold = _thresholds()
    comentarios = [normalizar(x) for x in (configuracao.STORE.get("comentarios_funcionario", default=[]) or [])]
    lancamentos = []
    dia_atual = mes_atual = None
    funcionario_atual = None
    cache_func = {}

    for linha_bruta in texto.splitlines():
        linha = RE_PREFIXO_WHATSAPP.sub("", linha_bruta.strip()).strip()
        if not linha or linha == "---":
            continue
        m = RE_DATA.match(linha)
        if m:
            dia_atual, mes_atual = int(m.group(1)), int(m.group(2))
            funcionario_atual = None
            continue
        m = RE_PRODUTO.match(linha)
        if m and dia_atual is not None:
            lancamentos.append({"dia": dia_atual, "mes": mes_atual, "funcionario_bruto": funcionario_atual,
                                "produto_bruto": m.group(2).strip(), "quantidade": int(m.group(1))})
            continue
        nome_limpo = linha.strip("*").strip()
        if not nome_limpo:
            continue
        nome_norm = normalizar(nome_limpo)
        if any(p and p in nome_norm for p in comentarios):
            continue
        if nome_norm not in cache_func:
            cache_func[nome_norm] = melhor_match_lista(nome_limpo, candidatos)
        _, score = cache_func[nome_norm]
        if score >= funcionario_threshold:
            funcionario_atual = nome_limpo
    return lancamentos


def processar_lancamentos(lancamentos, caminho=None, nomes_funcionarios=None, catalogos=None, dia_cols=None):
    """Processa lançamentos no catálogo REAL de cada funcionário.

    No Individual, primeiro tentamos o nome recebido diretamente. Isso é importante
    porque a planilha mensal possui linhas específicas como "Ninho Gota Tramado".
    Regras de Cadeiras/Mesas só entram como fallback quando o nome bruto não encontra
    nada confiável e o destino da regra também existe de fato naquela aba.
    """
    if catalogos is None or dia_cols is None:
        nomes_funcionarios, catalogos, dia_cols = carregar_estrutura_mensal(caminho)
    nomes_funcionarios = nomes_funcionarios or list(catalogos)
    candidatos_func = [(n, normalizar(n), normalizar(n).split()) for n in nomes_funcionarios]
    cache_func, cache_prod = {}, {}
    resultados = []
    _, revisar, _ = _thresholds()

    for lanc in lancamentos:
        excluido, motivo_exclusao = configuracao.exclusao_permanente(lanc["produto_bruto"])
        func_bruto = lanc["funcionario_bruto"] or ""
        norm_func = normalizar(func_bruto)
        if norm_func not in cache_func:
            cache_func[norm_func] = melhor_match_lista(func_bruto, candidatos_func)
        aba_match, score_func = cache_func[norm_func]

        produto_row = produto_nome = None
        score_prod = 0.0
        nota = motivo_exclusao if excluido else None

        if aba_match and not excluido:
            cat = catalogos.get(aba_match, [])
            itens = [(it["row"], it["name"], it["norm"], it["tokens"]) for it in cat]
            chave_bruta = (aba_match, "RAW", _normalizar_produto_barra(lanc["produto_bruto"]))
            if chave_bruta not in cache_prod:
                cache_prod[chave_bruta] = melhor_match_produto_lista(lanc["produto_bruto"], itens)
            row_match, score_prod = cache_prod[chave_bruta]

            # Fallback de regra: somente se o bruto ficou pendente e a regra apontar
            # para algo que realmente exista/tenha bom match dentro da aba atual.
            if score_prod < revisar:
                nome_regra, nota_regra = configuracao.aplicar_regras_produto(lanc["produto_bruto"])
                if nome_regra and nome_regra != "__SEM_CATALOGO__" and normalizar(nome_regra) != normalizar(lanc["produto_bruto"]):
                    chave_regra = (aba_match, "REGRA", _normalizar_produto_barra(nome_regra))
                    if chave_regra not in cache_prod:
                        cache_prod[chave_regra] = melhor_match_produto_lista(nome_regra, itens)
                    row_regra, score_regra = cache_prod[chave_regra]
                    if row_regra is not None and score_regra >= revisar and score_regra > score_prod + 8:
                        row_match, score_prod = row_regra, score_regra
                        nota = nota_regra

            if row_match is not None:
                produto_row = row_match
                produto_nome = next((it["name"] for it in cat if it["row"] == row_match), None)

        resultados.append({
            **lanc, "aba": aba_match, "score_func": round(score_func, 1), "produto_row": produto_row,
            "produto_nome": produto_nome, "score_prod": round(score_prod, 1), "manter": not excluido,
            "nota_regra": nota,
        })
    return resultados, catalogos, dia_cols


def agregar(resultados):
    """Agrega apenas reconhecimentos confiáveis ou escolhas manuais (score 100)."""
    _, revisar, funcionario_threshold = _thresholds()
    lancar = {}
    for r in resultados:
        if not r.get("manter") or not r.get("aba") or r.get("produto_row") is None:
            continue
        if r.get("score_func", 0) < funcionario_threshold or r.get("score_prod", 0) < revisar:
            continue
        chave = (r["aba"], r["produto_row"], r["dia"])
        lancar[chave] = lancar.get(chave, 0) + r["quantidade"]
    return lancar


def salvar_planilha(caminho, dia_cols, lancar, somar_existente, pasta_saida):
    wb = openpyxl.load_workbook(caminho, data_only=False, keep_vba=str(caminho).lower().endswith(".xlsm"))
    try:
        for (aba, row, dia), qtd in lancar.items():
            if aba not in wb.sheetnames:
                continue
            col = dia_cols.get(aba, {}).get(dia)
            if col is None:
                continue
            celula = wb[aba].cell(row=row, column=col)
            atual = celula.value
            celula.value = atual + qtd if somar_existente and isinstance(atual, (int, float)) else qtd
        destino = tema.caminho_saida_sugerido(caminho, pasta_saida)
        wb.save(destino)
        return destino
    finally:
        wb.close()


class ComboboxEditor:
    def __init__(self, tree, column, on_commit):
        self.tree, self.column, self.on_commit = tree, column, on_commit
        self.combo = self.iid = None

    def abrir(self, iid, valores, valor_atual):
        self.fechar(); self.iid = iid
        bbox = self.tree.bbox(iid, self.column)
        if not bbox: return
        x, y, w, h = bbox
        self.combo = ttk.Combobox(self.tree, values=valores, state="normal")
        self.combo.set(valor_atual or "")
        self.combo.place(x=x, y=y, width=max(w, 230), height=h)
        self.combo.focus_set(); self.combo.selection_range(0, tk.END)
        self.combo.bind("<<ComboboxSelected>>", self._confirmar)
        self.combo.bind("<Return>", self._confirmar)
        self.combo.bind("<Escape>", lambda e: self.fechar())
        self.combo.bind("<FocusOut>", self._confirmar)

    def _confirmar(self, _e=None):
        if self.combo is None: return
        valor, iid = self.combo.get().strip(), self.iid
        self.fechar()
        if valor: self.on_commit(iid, valor)

    def fechar(self):
        if self.combo is not None:
            self.combo.destroy()
        self.combo = self.iid = None



class SeletorBuscaDialog(tk.Toplevel):
    """Janela pesquisável no mesmo estilo do seletor de Cadeiras/Mesas."""

    def __init__(self, parent, titulo, itens, rotulo="Item", selecionado="", modo_produto=False):
        super().__init__(parent)
        self.title(titulo)
        self.geometry("700x520")
        self.minsize(520, 380)
        self.transient(parent)
        self.grab_set()
        self.resultado = None
        self.itens = sorted({str(x).strip() for x in (itens or []) if str(x).strip()}, key=normalizar)
        self.var_busca = tk.StringVar()
        self._selecionado = selecionado
        self.modo_produto = bool(modo_produto)

        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(2, weight=1)

        ttk.Label(root, text=f"Pesquisar {rotulo.lower()}", style="Section.TLabel").grid(row=0, column=0, sticky="w")
        ent = ttk.Entry(root, textvariable=self.var_busca)
        ent.grid(row=1, column=0, sticky="ew", pady=(6, 8))
        ent.focus_set()

        self.tree = ttk.Treeview(root, columns=("item",), show="headings", selectmode="browse")
        self.tree.heading("item", text=rotulo)
        self.tree.column("item", width=620, anchor="w")
        y = ttk.Scrollbar(root, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=y.set)
        self.tree.grid(row=2, column=0, sticky="nsew")
        y.grid(row=2, column=1, sticky="ns")

        botoes = ttk.Frame(root)
        botoes.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ttk.Button(botoes, text="Cancelar", command=self.destroy).pack(side="right")
        ttk.Button(botoes, text="Usar selecionado", style="Accent.TButton", command=self._confirmar).pack(side="right", padx=(0, 8))

        self.var_busca.trace_add("write", lambda *_: self._preencher())
        self.tree.bind("<Double-1>", lambda e: self._confirmar())
        self.tree.bind("<Return>", lambda e: self._confirmar())
        self.bind("<Escape>", lambda e: self.destroy())
        self._preencher()

    def _score_busca(self, termo, item):
        if not termo:
            return 100.0
        n = normalizar(item)
        if termo in n:
            return 130.0 - max(0, len(n) - len(termo)) * 0.04
        if self.modo_produto:
            # A busca manual usa o MESMO matcher inteligente do reconhecimento.
            return pontuar_nome_produto(self.var_busca.get(), item)
        return max(fuzz.WRatio(termo, n), fuzz.token_set_ratio(termo, n))

    def _preencher(self):
        termo = normalizar(self.var_busca.get())
        atual = self.resultado or self._selecionado
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        avaliados = [(self._score_busca(termo, item), item) for item in self.itens]
        if termo:
            avaliados = [x for x in avaliados if x[0] >= 42]
            avaliados.sort(key=lambda x: (-x[0], normalizar(x[1])))
        else:
            avaliados.sort(key=lambda x: normalizar(x[1]))
        selecionar = None
        for idx, (_score, item) in enumerate(avaliados):
            iid = str(idx)
            self.tree.insert("", "end", iid=iid, values=(item,))
            if atual and normalizar(item) == normalizar(atual):
                selecionar = iid
        if selecionar:
            self.tree.selection_set(selecionar)
            self.tree.see(selecionar)
        elif avaliados:
            self.tree.selection_set("0")

    def _confirmar(self):
        sel = self.tree.selection()
        if not sel:
            return
        self.resultado = str(self.tree.item(sel[0], "values")[0])
        self.destroy()


class LancadorIndividualFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        hist = configuracao.STORE.get("historico", default={}) or {}
        self.caminho_planilha = tk.StringVar(value=hist.get("individual_arquivo", ""))
        self.somar_existente = tk.BooleanVar(value=bool(configuracao.STORE.get("preferencias", "somar_existente", default=True)))
        self.nomes_funcionarios, self.resultados, self.catalogos, self.dia_cols = [], [], {}, {}
        self._ocupado = False
        self._build()
        self.bind_all("<<ConfiguracaoAlterada>>", self._config_changed, add="+")

    def _build(self):
        cab = ttk.Frame(self, padding=(16, 12, 16, 8)); cab.pack(fill="x")
        ttk.Label(cab, text="Lançador Mensal — por Funcionário", style="Header.TLabel").pack(anchor="w")
        ttk.Label(cab, text="Reconhece funcionário e produto, permite corrigir qualquer linha e grava uma cópia mensal atualizada.", style="SubHeader.TLabel").pack(anchor="w", pady=(2, 0))

        acoes = ttk.Frame(cab); acoes.pack(fill="x", pady=(9, 0))
        self.btn_processar = ttk.Button(acoes, text="1  ▶  Analisar", style="PrimaryAction.TButton", command=self.processar)
        self.btn_processar.pack(side="left")
        self.btn_salvar = ttk.Button(acoes, text="2  ✓  Lançar na planilha", style="LaunchAction.TButton", command=self.salvar, state="disabled")
        self.btn_salvar.pack(side="left", padx=(7, 0))
        ttk.Checkbutton(acoes, text="Somar ao existente", variable=self.somar_existente).pack(side="right")

        self.lbl_status = ttk.Label(self, text="Pronto.", style="Status.TLabel")
        self.lbl_status.pack(side="bottom", fill="x", padx=12, pady=(0, 8))

        rolavel = tema.ScrollableArea(self)
        rolavel.pack(fill="both", expand=True)
        corpo = rolavel.content

        arq = ttk.LabelFrame(corpo, text="Planilha", padding=10); arq.pack(fill="x", padx=12, pady=(0, 8)); arq.columnconfigure(0, weight=1)
        ttk.Label(arq, text="Planilha mensal por funcionário", style="Section.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 3))
        ttk.Entry(arq, textvariable=self.caminho_planilha).grid(row=1, column=0, sticky="ew")
        ttk.Button(arq, text="Selecionar…", style="Small.TButton", command=self._selecionar_arquivo).grid(row=1, column=1, padx=(6, 0))

        paned = ttk.Panedwindow(corpo, orient="vertical"); paned.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        entrada, preview = ttk.Frame(paned), ttk.Frame(paned)
        paned.add(entrada, weight=3); paned.add(preview, weight=7)
        top = ttk.Frame(entrada); top.pack(fill="x", pady=(0, 5))
        ttk.Label(top, text="Mensagem / relatório do WhatsApp", style="Section.TLabel").pack(side="left")
        ttk.Button(top, text="Limpar", style="Small.TButton", command=lambda: self.txt_entrada.delete("1.0", tk.END)).pack(side="right")
        self.txt_entrada = scrolledtext.ScrolledText(entrada, wrap="word", height=9, font=("Consolas", 10)); tema.estilizar_texto(self.txt_entrada); self.txt_entrada.pack(fill="both", expand=True)
        preview.columnconfigure(0, weight=1); preview.rowconfigure(1, weight=1)
        ttk.Label(preview, text="Prévia editável — duplo clique em Funcionário, Produto ou Manter", style="Section.TLabel").grid(row=0, column=0, sticky="w", pady=(3, 5))
        cols = ("data", "func_bruto", "funcionario", "score_f", "produto_bruto", "produto", "score_p", "qtd", "manter", "nota")
        self.tree = ttk.Treeview(preview, columns=cols, show="headings", selectmode="extended")
        defs = [("data", "Data", 70), ("func_bruto", "Funcionário recebido", 130), ("funcionario", "Funcionário", 135), ("score_f", "%", 45), ("produto_bruto", "Produto recebido", 180), ("produto", "Produto", 210), ("score_p", "%", 45), ("qtd", "Qtd", 45), ("manter", "Manter", 60), ("nota", "Observação", 210)]
        for c, t, w in defs: self.tree.heading(c, text=t); self.tree.column(c, width=w, anchor="w")
        y = ttk.Scrollbar(preview, orient="vertical", command=self.tree.yview); x = ttk.Scrollbar(preview, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)
        self.tree.grid(row=1, column=0, sticky="nsew"); y.grid(row=1, column=1, sticky="ns"); x.grid(row=2, column=0, sticky="ew")
        for tag, (bg, fg) in tema.TAG_CORES.items(): self.tree.tag_configure(tag, background=bg, foreground=fg)
        self.tree.bind("<Double-1>", self._on_double_click); self.tree.bind("<Button-3>", self._menu_contexto)

        self.menu = tk.Menu(self, tearoff=False)
        self.menu.add_command(label="Alterar funcionário…", command=self._alterar_funcionario_menu)
        self.menu.add_command(label="Alterar produto…", command=self._alterar_produto_menu)
        self.menu.add_separator()
        self.menu.add_command(label="Alternar Manter", command=self._alternar_manter)
        self.menu.add_command(label="Excluir sempre este produto", command=self._excluir_sempre)
        self.menu.add_command(label="Restaurar nesta execução", command=self._restaurar_selecionados)

    def _selecionar_arquivo(self):
        atual = self.caminho_planilha.get().strip()
        initial = os.path.dirname(atual) if atual and os.path.isdir(os.path.dirname(atual)) else None
        path = filedialog.askopenfilename(
            title="Selecionar planilha mensal por funcionário",
            initialdir=initial,
            filetypes=[("Planilhas Excel", "*.xlsx *.xlsm"), ("Todos os arquivos", "*.*")],
        )
        if path:
            self.caminho_planilha.set(path)
            configuracao.salvar_historico("individual_arquivo", path)

    def _set_busy(self, busy, texto=None):
        self._ocupado = busy
        self.btn_processar.config(state="disabled" if busy else "normal")
        if busy: self.btn_salvar.config(state="disabled")
        if texto: self.lbl_status.config(text=texto)

    def _config_changed(self, _e=None):
        if not self.resultados:
            self.somar_existente.set(bool(configuracao.STORE.get("preferencias", "somar_existente", default=True)))

    def processar(self):
        caminho = self.caminho_planilha.get().strip()
        ok, erro = tema.validar_caminho_planilha(caminho)
        if not ok: messagebox.showerror("Planilha inválida", erro); return
        texto = self.txt_entrada.get("1.0", "end")
        if not texto.strip(): messagebox.showwarning("Nada colado", "Cole a mensagem antes de processar."); return
        self._set_busy(True, "Carregando a estrutura mensal uma única vez...")

        def work(report):
            nomes, cats, dias = carregar_estrutura_mensal(caminho, lambda p, t: report(("load", p, t)))
            if not nomes: raise ValueError("Nenhuma aba de funcionário com layout de dias foi encontrada.")
            lanc = parse_texto(texto, nomes)
            if not lanc: raise ValueError("Não consegui identificar nenhum lançamento no texto colado.")
            res, cats, dias = processar_lancamentos(lanc, nomes_funcionarios=nomes, catalogos=cats, dia_cols=dias)
            return nomes, cats, dias, res
        def progress(v):
            _, p, t = v; self.lbl_status.config(text=f"Carregando funcionários... {p}/{t}")
        def done(v):
            self.nomes_funcionarios, self.catalogos, self.dia_cols, self.resultados = v
            configuracao.salvar_historico("individual_arquivo", caminho)
            self._set_busy(False); self._repintar_tabela()
        def fail(exc):
            self._set_busy(False, "Falha ao processar."); messagebox.showerror("Erro ao processar", str(exc))
        tema.executar_assincrono(self, work, done, fail, progress)

    def _tag_linha(self, r):
        aceitar, revisar, _ = _thresholds()
        if not r.get("manter"): return "excluido"
        if not r.get("aba") or r.get("produto_row") is None: return "pendente"
        pior = min(r.get("score_func", 0), r.get("score_prod", 0))
        if pior < revisar: return "pendente"
        if pior < aceitar: return "revisar"
        return "ok"

    def _valores(self, r):
        return (f"{r['dia']:02d}/{r['mes']:02d}", r.get("funcionario_bruto") or "-", r.get("aba") or "(não identificado)",
                r.get("score_func", 0), r["produto_bruto"], r.get("produto_nome") or "(não identificado)", r.get("score_prod", 0),
                r["quantidade"], "Sim" if r.get("manter") else "Não", r.get("nota_regra") or "")

    def _repintar_tabela(self):
        for iid in self.tree.get_children(): self.tree.delete(iid)
        for idx, r in enumerate(self.resultados): self.tree.insert("", "end", iid=str(idx), tags=(self._tag_linha(r),), values=self._valores(r))
        lancar = agregar(self.resultados)
        _, revisar, funcionario_threshold = _thresholds()
        sem_func = sum(1 for r in self.resultados if r.get("manter") and (not r.get("aba") or r.get("score_func", 0) < funcionario_threshold))
        sem_prod = sum(1 for r in self.resultados if r.get("manter") and r.get("aba") and (r.get("produto_row") is None or r.get("score_prod", 0) < revisar))
        excl = sum(1 for r in self.resultados if not r.get("manter"))
        self.btn_salvar.config(state="normal" if lancar and not self._ocupado else "disabled")
        self.lbl_status.config(text=f"{len(self.resultados)} linha(s) • {sum(lancar.values())} unidade(s) prontas • {sem_func} funcionário(s) pendente(s) • {sem_prod} produto(s) pendente(s) • {excl} excluída(s)")

    def _on_double_click(self, event):
        iid, col = self.tree.identify_row(event.y), self.tree.identify_column(event.x)
        if not iid:
            return
        r = self.resultados[int(iid)]
        if col == "#3":
            d = SeletorBuscaDialog(self, "Escolher funcionário", self.nomes_funcionarios,
                                   rotulo="Funcionário / aba", selecionado=r.get("aba") or "")
            self.wait_window(d)
            if d.resultado:
                self._trocar_funcionario(iid, d.resultado)
        elif col == "#6":
            aba = r.get("aba")
            if not aba:
                messagebox.showinfo("Funcionário primeiro", "Escolha o funcionário antes do produto.")
                return
            vals = [it["name"] for it in self.catalogos.get(aba, [])]
            d = SeletorBuscaDialog(self, f"Escolher produto — {aba}", vals,
                                   rotulo="Produto", selecionado=r.get("produto_nome") or "", modo_produto=True)
            self.wait_window(d)
            if d.resultado:
                self._trocar_produto(iid, d.resultado)
        elif col == "#9":
            r["manter"] = not r.get("manter")
            r["nota_regra"] = "alterado manualmente"
            self._repintar_tabela()

    def _menu_contexto(self, event):
        iid = self.tree.identify_row(event.y)
        if iid and iid not in self.tree.selection(): self.tree.selection_set(iid)
        if iid:
            try: self.menu.tk_popup(event.x_root, event.y_root)
            finally: self.menu.grab_release()

    def _selected_indices(self): return [int(i) for i in self.tree.selection()]

    def _trocar_funcionario(self, iid, novo_nome):
        r = self.resultados[int(iid)]
        if novo_nome not in self.nomes_funcionarios:
            cand = [(n, normalizar(n), normalizar(n).split()) for n in self.nomes_funcionarios]
            novo_nome, _ = melhor_match_lista(novo_nome, cand)
            if not novo_nome: return
        r["aba"], r["score_func"], r["manter"], r["nota_regra"] = novo_nome, 100.0, True, "funcionário definido manualmente"
        cat = self.catalogos.get(novo_nome, [])
        itens = [(it["row"], it["name"], it["norm"], it["tokens"]) for it in cat]
        row, score = melhor_match_produto_lista(r["produto_bruto"], itens)
        r["produto_row"], r["score_prod"] = row, round(score, 1)
        r["produto_nome"] = next((it["name"] for it in cat if it["row"] == row), None)
        self._repintar_tabela()

    def _trocar_produto(self, iid, novo_nome):
        r = self.resultados[int(iid)]
        cat = self.catalogos.get(r.get("aba"), [])
        item = next((it for it in cat if it["name"] == novo_nome), None)
        if item is None:
            row, _ = melhor_match_produto_lista(novo_nome, [(it["row"], it["name"], it["norm"], it["tokens"]) for it in cat])
            item = next((it for it in cat if it["row"] == row), None)
        if item:
            r.update(produto_row=item["row"], produto_nome=item["name"], score_prod=100.0, manter=True, nota_regra="produto definido manualmente")
            self._repintar_tabela()

    def _alterar_funcionario_menu(self):
        ids = self.tree.selection()
        if not ids:
            return
        iid = ids[0]
        r = self.resultados[int(iid)]
        d = SeletorBuscaDialog(self, "Escolher funcionário", self.nomes_funcionarios,
                               rotulo="Funcionário / aba", selecionado=r.get("aba") or "")
        self.wait_window(d)
        if d.resultado:
            self._trocar_funcionario(iid, d.resultado)

    def _alterar_produto_menu(self):
        ids = self.tree.selection()
        if not ids:
            return
        iid = ids[0]
        r = self.resultados[int(iid)]
        aba = r.get("aba")
        if not aba:
            messagebox.showinfo("Funcionário primeiro", "Escolha o funcionário antes do produto.")
            return
        vals = [it["name"] for it in self.catalogos.get(aba, [])]
        d = SeletorBuscaDialog(self, f"Escolher produto — {aba}", vals,
                               rotulo="Produto", selecionado=r.get("produto_nome") or "", modo_produto=True)
        self.wait_window(d)
        if d.resultado:
            self._trocar_produto(iid, d.resultado)

    def _alternar_manter(self):
        for idx in self._selected_indices():
            r = self.resultados[idx]; r["manter"] = not r.get("manter"); r["nota_regra"] = "alterado manualmente"
        self._repintar_tabela()

    def _excluir_sempre(self):
        count = 0
        for idx in self._selected_indices():
            r = self.resultados[idx]
            if configuracao.adicionar_exclusao(r["produto_bruto"], "exato"): count += 1
            r["manter"] = False; r["nota_regra"] = "exclusão permanente (nome exato)"
        self.event_generate("<<ConfiguracaoAlterada>>", when="tail")
        self._repintar_tabela()
        if count: self.lbl_status.config(text=f"{count} produto(s) adicionado(s) às exclusões permanentes.")

    def _restaurar_selecionados(self):
        for idx in self._selected_indices():
            r = self.resultados[idx]; r["manter"] = True; r["nota_regra"] = "restaurado manualmente nesta execução"
            if r.get("aba") and r.get("produto_row") is None:
                cat = self.catalogos.get(r["aba"], [])
                row, score = melhor_match_produto_lista(r["produto_bruto"], [(it["row"], it["name"], it["norm"], it["tokens"]) for it in cat])
                r["produto_row"], r["score_prod"] = row, round(score, 1)
                r["produto_nome"] = next((it["name"] for it in cat if it["row"] == row), None)
        self._repintar_tabela()

    def salvar(self):
        lancar = agregar(self.resultados)
        if not lancar: return
        total = sum(lancar.values())
        confirmar = bool(configuracao.STORE.get("preferencias", "confirmar_salvamento", default=True))
        _, revisar, funcionario_threshold = _thresholds()
        sem = sum(
            1 for r in self.resultados
            if r.get("manter") and (
                not r.get("aba") or r.get("produto_row") is None
                or r.get("score_func", 0) < funcionario_threshold
                or r.get("score_prod", 0) < revisar
            )
        )
        aviso = f"\n\n{sem} linha(s) pendente(s)/de baixa confiança não serão lançadas." if sem else ""
        if confirmar and not messagebox.askyesno("Confirmar lançamento", f"Serão lançadas {total} unidade(s).{aviso}\n\nContinuar?"): return
        initial = configuracao.STORE.get("historico", "pasta_saida", default="") or ""
        pasta = filedialog.askdirectory(title="Pasta para salvar a planilha", initialdir=initial if os.path.isdir(initial) else None)
        if not pasta: return
        destino = tema.caminho_saida_sugerido(self.caminho_planilha.get(), pasta)
        if not tema.confirmar_sobrescrita(destino): return
        self._set_busy(True, "Salvando a cópia mensal em segundo plano...")
        def work(report): return salvar_planilha(self.caminho_planilha.get(), self.dia_cols, lancar, self.somar_existente.get(), pasta)
        def done(path):
            self._set_busy(False); self._repintar_tabela(); configuracao.salvar_historico("pasta_saida", pasta); messagebox.showinfo("Concluído", f"Planilha salva em:\n{path}")
        def fail(exc): self._set_busy(False, "Falha ao salvar."); messagebox.showerror("Erro ao salvar", str(exc))
        tema.executar_assincrono(self, work, done, fail)


if __name__ == "__main__":
    root = tk.Tk(); root.title("Lançador Mensal — por Funcionário"); root.geometry("1280x820"); root.minsize(980, 650)
    tema.aplicar_tema_escuro(root); LancadorIndividualFrame(root).pack(fill="both", expand=True); root.mainloop()
