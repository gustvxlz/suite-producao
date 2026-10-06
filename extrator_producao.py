"""Extrator de Produção otimizado.

Abre o arquivo Excel apenas uma vez por extração, varre cada aba selecionada
uma única vez e não depende de pandas. Isso reduz bastante tempo de abertura,
RAM e leituras de disco em PCs básicos.
"""

from __future__ import annotations

import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

import openpyxl

try:
    from . import tema, configuracao
except ImportError:
    import tema
    import configuracao


_SHEET_CACHE = {}


def _assinatura_arquivo(path):
    st = os.stat(path)
    return os.path.abspath(path), st.st_mtime_ns, st.st_size


def listar_abas(file_path):
    """Lista abas sem carregar células; cacheia por mtime/tamanho do arquivo."""
    sig = _assinatura_arquivo(file_path)
    cached = _SHEET_CACHE.get(sig)
    if cached is not None:
        return list(cached)
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets = [s.strip() for s in wb.sheetnames if "TOTAL" not in s.upper()]
    finally:
        wb.close()
    _SHEET_CACHE.clear()  # basta o último arquivo; evita crescimento infinito
    _SHEET_CACHE[sig] = tuple(sheets)
    return sheets


def _encontrar_layout(ws):
    """Localiza coluna PRODUTOS e linha de cabeçalho dos dias de forma tolerante."""
    produto_col = None
    header_row = None
    max_scan_rows = min(ws.max_row or 1, 15)
    max_scan_cols = min(ws.max_column or 1, 80)

    for row in ws.iter_rows(min_row=1, max_row=max_scan_rows, max_col=max_scan_cols):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().upper() == "PRODUTOS":
                produto_col = cell.column
                header_row = cell.row
                break
        if produto_col:
            break
    if not produto_col:
        return None, None, {}

    melhor_linha = None
    melhor_mapa = {}
    # O layout original usa a primeira linha de dados como números dos dias,
    # mas procuramos algumas linhas para tolerar planilhas levemente diferentes.
    for r in range(header_row + 1, min(header_row + 5, ws.max_row or header_row + 1) + 1):
        mapa = {}
        for c in range(1, max_scan_cols + 1):
            if c == produto_col:
                continue
            v = ws.cell(row=r, column=c).value
            try:
                fv = float(v)
                if fv.is_integer() and 1 <= int(fv) <= 31:
                    mapa[int(fv)] = c
            except (TypeError, ValueError):
                continue
        if len(mapa) > len(melhor_mapa):
            melhor_linha, melhor_mapa = r, mapa
    return produto_col, melhor_linha, melhor_mapa


def _formatar_qtd(qty):
    try:
        fq = float(qty)
        if fq.is_integer():
            return str(int(fq))
    except (TypeError, ValueError):
        pass
    return str(qty)


def extract_production_data(file_path, professionals, start_day, end_day, month, progresso=None):
    """Extrai produção com uma única abertura do workbook.

    `progresso(passo, total)` é opcional e mede abas processadas.
    """
    if end_day < start_day:
        raise ValueError("O dia final não pode ser menor que o dia inicial.")
    dias = list(range(start_day, end_day + 1))
    por_dia = {d: [] for d in dias}
    wb = openpyxl.load_workbook(file_path, read_only=False, data_only=True)
    try:
        total = max(1, len(professionals))
        for passo, sheet in enumerate(professionals, start=1):
            if sheet not in wb.sheetnames:
                if progresso:
                    progresso(passo, total)
                continue
            ws = wb[sheet]
            produto_col, linha_dias, dia_cols = _encontrar_layout(ws)
            if not produto_col or not linha_dias:
                if progresso:
                    progresso(passo, total)
                continue

            cols_interesse = {d: dia_cols.get(d) for d in dias if dia_cols.get(d)}
            if not cols_interesse:
                if progresso:
                    progresso(passo, total)
                continue

            # Varre as linhas de produto apenas uma vez e alimenta todos os dias.
            for r in range(linha_dias + 1, (ws.max_row or linha_dias) + 1):
                product = ws.cell(row=r, column=produto_col).value
                if product is None:
                    continue
                product_text = str(product).strip()
                if not product_text or product_text.upper() == "TOTAIS":
                    continue
                if product_text == "0":
                    product_text = "Cade Happy hour"
                for dia, col in cols_interesse.items():
                    qty = ws.cell(row=r, column=col).value
                    if qty is None or qty == 0:
                        continue
                    # Proteção equivalente ao código antigo contra linhas de cabeçalho.
                    try:
                        if float(product) == float(dia):
                            continue
                    except (TypeError, ValueError):
                        pass
                    por_dia[dia].append((sheet, product_text, qty))
            if progresso:
                progresso(passo, total)
    finally:
        wb.close()

    output_lines = []
    for dia in dias:
        output_lines.append(f"{dia:02d}/{month:02d}")
        ultimo_sheet = None
        for sheet, product, qty in por_dia[dia]:
            if sheet != ultimo_sheet:
                output_lines.append(f"\n**{sheet}**")
                ultimo_sheet = sheet
            output_lines.append(f"{_formatar_qtd(qty)} {product}")
        output_lines.append("\n---\n")
    return "\n".join(output_lines)


class ExtratorFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        hoje = __import__("datetime").datetime.now()
        historico = configuracao.STORE.get("historico", default={}) or {}
        self.file_path = tk.StringVar(value=historico.get("extrator_arquivo", ""))
        self.start_day = tk.StringVar(value="1")
        self.end_day = tk.StringVar(value=str(hoje.day))
        self.month = tk.StringVar(value=str(hoje.month))
        self.var_busca = tk.StringVar()
        self.var_grupo = tk.StringVar()
        self._todas_abas = []
        self._ocupado = False
        self._build_layout()
        self.bind_all("<<ConfiguracaoAlterada>>", self._on_config_change, add="+")
        self._refresh_group_combo()

    def _build_layout(self):
        cab = ttk.Frame(self, padding=(16, 12, 16, 8))
        cab.pack(fill="x")
        ttk.Label(cab, text="Extrator de Produção", style="Header.TLabel").pack(anchor="w")
        ttk.Label(cab, text="Leitura única da planilha, seleção rápida de equipes e geração de relatório por dia.", style="SubHeader.TLabel").pack(anchor="w", pady=(2, 0))

        # Ações principais no topo, compactas e dentro da mesma paleta do app.
        acoes = ttk.Frame(cab)
        acoes.pack(fill="x", pady=(9, 0))
        self.btn_extrair = ttk.Button(acoes, text="▶  Gerar relatório", style="PrimaryAction.TButton",
                                      command=self._extract, state="disabled")
        self.btn_extrair.pack(side="left")
        self.btn_copiar = ttk.Button(acoes, text="Copiar resultado", style="Small.TButton",
                                     command=self._copy_output, state="disabled")
        self.btn_copiar.pack(side="left", padx=(7, 0))
        self.btn_salvar = ttk.Button(acoes, text="Salvar .txt", style="Small.TButton",
                                     command=self._save_output, state="disabled")
        self.btn_salvar.pack(side="left", padx=(7, 0))
        self.progresso = ttk.Progressbar(cab, style="Accent.Horizontal.TProgressbar", maximum=100)
        self.progresso.pack(fill="x", pady=(7, 0))

        self.lbl_status = ttk.Label(self, text="Selecione uma planilha para começar.", style="Status.TLabel")
        self.lbl_status.pack(side="bottom", fill="x", padx=12, pady=(0, 8))

        rolavel = tema.ScrollableArea(self)
        rolavel.pack(fill="both", expand=True)
        corpo = rolavel.content

        paned = ttk.Panedwindow(corpo, orient="horizontal")
        paned.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        left = ttk.Frame(paned, padding=(4, 4, 10, 4))
        right = ttk.Frame(paned, padding=(10, 4, 4, 4))
        paned.add(left, weight=4)
        paned.add(right, weight=6)

        frm_file = ttk.LabelFrame(left, text="Arquivo", padding=10)
        frm_file.pack(fill="x", pady=(0, 8))
        frm_file.columnconfigure(0, weight=1)
        ttk.Entry(frm_file, textvariable=self.file_path).grid(row=0, column=0, sticky="ew")
        ttk.Button(frm_file, text="Procurar", command=self._browse_file).grid(row=0, column=1, padx=(6, 0))
        ttk.Button(frm_file, text="Recarregar abas", style="Small.TButton", command=self._load_sheets).grid(row=1, column=0, columnspan=2, sticky="ew", pady=(7, 0))

        periodo = ttk.LabelFrame(left, text="Período", padding=10)
        periodo.pack(fill="x", pady=(0, 8))
        for i, (rot, var, fim) in enumerate((("Inicial", self.start_day, 31), ("Final", self.end_day, 31), ("Mês", self.month, 12))):
            ttk.Label(periodo, text=rot).grid(row=0, column=i, sticky="w", padx=(0 if i == 0 else 10, 0))
            ttk.Spinbox(periodo, from_=1, to=fim, textvariable=var, width=6, justify="center").grid(row=1, column=i, sticky="w", padx=(0 if i == 0 else 10, 0), pady=(3, 0))

        prof = ttk.LabelFrame(left, text="Profissionais", padding=10)
        prof.pack(fill="both", expand=True)
        top = ttk.Frame(prof)
        top.pack(fill="x", pady=(0, 7))
        self.cmb_grupo = ttk.Combobox(top, textvariable=self.var_grupo, state="readonly", width=22)
        self.cmb_grupo.pack(side="left", fill="x", expand=True)
        ttk.Button(top, text="Selecionar grupo", style="Small.TButton", command=self._select_current_group).pack(side="left", padx=(6, 0))

        busca = ttk.Entry(prof, textvariable=self.var_busca)
        busca.pack(fill="x", pady=(0, 7))
        self.var_busca.trace_add("write", lambda *_: self._aplicar_filtro())

        lc = ttk.Frame(prof)
        lc.pack(fill="both", expand=True)
        scroll = ttk.Scrollbar(lc, orient="vertical")
        self.listbox = tk.Listbox(lc, selectmode=tk.EXTENDED, exportselection=False, yscrollcommand=scroll.set, height=12)
        tema.estilizar_listbox(self.listbox)
        scroll.config(command=self.listbox.yview)
        self.listbox.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        self.listbox.bind("<<ListboxSelect>>", lambda e: self._atualizar_estado())

        sel = ttk.Frame(prof)
        sel.pack(fill="x", pady=(7, 0))
        ttk.Button(sel, text="Todos", style="Small.TButton", command=self._select_all).pack(side="left")
        ttk.Button(sel, text="Limpar", style="Small.TButton", command=self._clear_selection).pack(side="left", padx=6)
        self.lbl_sel = ttk.Label(sel, text="0 selecionados", style="Muted.TLabel")
        self.lbl_sel.pack(side="right")

        right.columnconfigure(0, weight=1)
        right.rowconfigure(1, weight=1)
        ttk.Label(right, text="Resultado", style="Section.TLabel").grid(row=0, column=0, sticky="w", pady=(0, 6))
        self.output = scrolledtext.ScrolledText(right, wrap="word", font=("Consolas", 10))
        tema.estilizar_texto(self.output)
        self.output.grid(row=1, column=0, sticky="nsew")

        if self.file_path.get() and os.path.isfile(self.file_path.get()):
            self.after(80, self._load_sheets)

    def _on_config_change(self, _e=None):
        self._refresh_group_combo()

    def _refresh_group_combo(self):
        grupos = configuracao.STORE.get("grupos", default={}) or {}
        nomes = list(grupos)
        self.cmb_grupo["values"] = nomes
        if self.var_grupo.get() not in nomes:
            self.var_grupo.set(nomes[0] if nomes else "")

    def _browse_file(self):
        path = filedialog.askopenfilename(title="Selecione a planilha", filetypes=[("Excel", "*.xlsx *.xlsm")])
        if path:
            self.file_path.set(path)
            configuracao.salvar_historico("extrator_arquivo", path)
            self._load_sheets()

    def _load_sheets(self):
        path = self.file_path.get().strip()
        ok, erro = tema.validar_caminho_planilha(path)
        if not ok:
            messagebox.showerror("Arquivo inválido", erro)
            return
        self.lbl_status.config(text="Lendo nomes das abas...")
        self._set_busy(True)

        def work(report):
            return listar_abas(path)

        def done(sheets):
            self._todas_abas = sheets
            self.var_busca.set("")
            self._aplicar_filtro()
            self._set_busy(False)
            self.lbl_status.config(text=f"{len(sheets)} profissional(is) carregado(s).")
            configuracao.salvar_historico("extrator_arquivo", path)

        def fail(exc):
            self._set_busy(False)
            self.lbl_status.config(text="Falha ao carregar a planilha.")
            messagebox.showerror("Erro ao ler Excel", str(exc))

        tema.executar_assincrono(self, work, done, fail)

    def _aplicar_filtro(self):
        termo = self.var_busca.get().strip().upper()
        selecionados = {self.listbox.get(i) for i in self.listbox.curselection()} if hasattr(self, "listbox") else set()
        self.listbox.delete(0, tk.END)
        for s in self._todas_abas:
            if not termo or termo in s.upper():
                self.listbox.insert(tk.END, s)
        for i, s in enumerate(self.listbox.get(0, tk.END)):
            if s in selecionados:
                self.listbox.selection_set(i)
        self._atualizar_estado()

    def _select_all(self):
        self.listbox.selection_set(0, tk.END)
        self._atualizar_estado()

    def _clear_selection(self):
        self.listbox.selection_clear(0, tk.END)
        self._atualizar_estado()

    def _select_current_group(self):
        grupo = self.var_grupo.get()
        grupos = configuracao.STORE.get("grupos", default={}) or {}
        alvos = {str(n).strip().upper() for n in grupos.get(grupo, [])}
        if not alvos:
            return
        # Grupo deve funcionar mesmo com filtro ativo: limpamos a busca para mostrar tudo.
        self.var_busca.set("")
        self.listbox.selection_clear(0, tk.END)
        for i, item in enumerate(self.listbox.get(0, tk.END)):
            if item.strip().upper() in alvos:
                self.listbox.selection_set(i)
        if not self.listbox.curselection():
            messagebox.showwarning("Grupo sem correspondências", f"Nenhum funcionário do grupo '{grupo}' existe nessa planilha.")
        self._atualizar_estado()

    def _atualizar_estado(self):
        n = len(self.listbox.curselection())
        self.lbl_sel.config(text=f"{n} selecionado(s)")
        self.btn_extrair.config(state="disabled" if self._ocupado or n == 0 else "normal")

    def _set_busy(self, busy):
        self._ocupado = busy
        self._atualizar_estado()

    def _extract(self):
        path = self.file_path.get().strip()
        ok, erro = tema.validar_caminho_planilha(path)
        if not ok:
            messagebox.showerror("Arquivo inválido", erro); return
        try:
            ini, fim, mes = int(self.start_day.get()), int(self.end_day.get()), int(self.month.get())
        except ValueError:
            messagebox.showerror("Período inválido", "Dia e mês precisam ser números."); return
        if not (1 <= ini <= 31 and 1 <= fim <= 31 and 1 <= mes <= 12) or ini > fim:
            messagebox.showerror("Período inválido", "Confira os dias inicial/final e o mês."); return
        indices = self.listbox.curselection()
        if not indices:
            return
        profissionais = [self.listbox.get(i) for i in indices]
        self._set_busy(True)
        self.btn_salvar.config(state="disabled")
        self.btn_copiar.config(state="disabled")
        self.progresso.config(value=0)
        self.lbl_status.config(text="Extraindo sem bloquear a interface...")

        def work(report):
            def prog(p, total):
                report((p, total))
            return extract_production_data(path, profissionais, ini, fim, mes, prog)

        def progress(info):
            p, total = info
            self.progresso.config(value=(p / max(total, 1)) * 100)
            self.lbl_status.config(text=f"Lendo profissionais... {p}/{total}")

        def done(result):
            self.output.delete("1.0", tk.END)
            self.output.insert("1.0", result)
            self.progresso.config(value=100)
            self._set_busy(False)
            self.btn_salvar.config(state="normal")
            self.btn_copiar.config(state="normal")
            self.lbl_status.config(text=f"Concluído — {len(profissionais)} profissional(is), dias {ini:02d} a {fim:02d}/{mes:02d}.")

        def fail(exc):
            self._set_busy(False)
            self.lbl_status.config(text="Falha na extração.")
            messagebox.showerror("Erro ao extrair", str(exc))

        tema.executar_assincrono(self, work, done, fail, progress)

    def _copy_output(self):
        content = self.output.get("1.0", tk.END).strip()
        if not content:
            return
        try:
            self.clipboard_clear()
            self.clipboard_append(content)
            self.update_idletasks()
            self.lbl_status.config(text="Resultado copiado — pronto para colar no WhatsApp ou onde quiser.")
        except tk.TclError as e:
            messagebox.showerror("Erro ao copiar", str(e))

    def _save_output(self):
        content = self.output.get("1.0", tk.END).strip()
        if not content:
            return
        pasta = configuracao.STORE.get("historico", "pasta_saida", default="") or ""
        path = filedialog.asksaveasfilename(title="Salvar relatório", defaultextension=".txt", filetypes=[("Texto", "*.txt")], initialdir=pasta if os.path.isdir(pasta) else None)
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write(content)
            configuracao.salvar_historico("pasta_saida", os.path.dirname(path))
            self.lbl_status.config(text=f"Relatório salvo em {path}")
        except OSError as e:
            messagebox.showerror("Erro ao salvar", str(e))


if __name__ == "__main__":
    root = tk.Tk()
    root.title("Extrator de Produção")
    root.geometry("1120x720")
    root.minsize(850, 600)
    tema.aplicar_tema_escuro(root)
    ExtratorFrame(root).pack(fill="both", expand=True)
    root.mainloop()
