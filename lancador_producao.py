"""Lançador de Produção — Cadeiras / Mesas, versão otimizada e configurável."""

from __future__ import annotations

import os
import re
import unicodedata
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

import openpyxl
from rapidfuzz import fuzz

try:
    from . import tema, configuracao
except ImportError:
    import tema
    import configuracao

SHEET_NAME = "PRODUÇÃO-DIA-MÊS-QUANT"
PRODUTO_COL = 3
DAY_HEADER_COL_START = 9
DAY_HEADER_COL_END = 31

# Mantidos por compatibilidade com scripts que importavam estas constantes.
ACCEPT_THRESHOLD = 85
REVIEW_THRESHOLD = 60


def _thresholds():
    aceitar, revisar, _ = configuracao.thresholds()
    return aceitar, revisar


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    texto = texto.upper()
    texto = re.sub(r"[^A-Z0-9 ]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _abrir_workbook(caminho, read_only=False, data_only=False):
    return openpyxl.load_workbook(
        caminho,
        read_only=read_only,
        data_only=data_only,
        keep_vba=str(caminho).lower().endswith(".xlsm"),
    )


def carregar_catalogo(caminho, tag):
    """Lê catálogo em uma única passagem, com workbook read-only."""
    wb = _abrir_workbook(caminho, read_only=True, data_only=False)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"A aba '{SHEET_NAME}' não foi encontrada em {os.path.basename(caminho)}")
        ws = wb[SHEET_NAME]
        header_row = None
        header_values = None
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 100), values_only=True), start=1):
            if len(row) >= PRODUTO_COL and row[PRODUTO_COL - 1] == "ESPECIFICAÇÃO":
                header_row = r
                header_values = row
                break
        if header_row is None:
            raise ValueError(f"Não encontrei o cabeçalho 'ESPECIFICAÇÃO' em {os.path.basename(caminho)}")

        dia_para_coluna = {}
        # O próprio cabeçalho já veio na iteração, evitando acessos célula a célula.
        for c in range(DAY_HEADER_COL_START, min(DAY_HEADER_COL_END, len(header_values)) + 1):
            v = header_values[c - 1]
            if isinstance(v, int) and 1 <= v <= 31:
                dia_para_coluna[v] = c
            elif isinstance(v, float) and v.is_integer() and 1 <= int(v) <= 31:
                dia_para_coluna[int(v)] = c

        catalogo = []
        excluir_nomes = {"PRODUTO", "ESPECIFICAÇÃO"}
        start = header_row + 2
        for r, row in enumerate(ws.iter_rows(min_row=start, values_only=True), start=start):
            nome = row[PRODUTO_COL - 1] if len(row) >= PRODUTO_COL else None
            c1 = row[0] if row else None
            if not isinstance(nome, str) or not nome.strip():
                continue
            nome = nome.strip()
            if nome in excluir_nomes or c1 in ("Nº", "ORD"):
                continue
            norm = normalizar(nome)
            tokens = norm.split()
            if not tokens:
                continue
            catalogo.append({"row": r, "name": nome, "norm": norm, "tokens": tokens, "tag": tag})
        return catalogo, dia_para_coluna
    finally:
        wb.close()



# Matching hierárquico de produtos, ajustado contra as planilhas reais de agosto/2026.
_TIPOS_PRODUTO = {
    "CADE": "CADE", "CADEIRA": "CADE", "CADEIRAS": "CADE",
    "MESA": "MESA", "MESAS": "MESA",
    "POLT": "POLT", "POLTRONA": "POLT", "POLTRONAS": "POLT",
    "BANQ": "BANQ", "BANQUETA": "BANQ", "BANQUETAS": "BANQ",
    "NAMO": "NAMO", "NAMORADEIRA": "NAMO", "NAMORADEIRAS": "NAMO",
    "BALANCO": "BALANCO", "BALANCOS": "BALANCO",
    "CHAISE": "CHAISE", "CHAISES": "CHAISE",
    "ESPR": "ESPR", "ESP": "ESPR", "ESPREGUICADEIRA": "ESPR", "ESPREGUICADEIRAS": "ESPR",
    "BISTRO": "BISTRO", "CACH": "CACH", "CACHEPO": "CACH", "CACHEPOS": "CACH",
    "PUFF": "PUFF", "ASSENTO": "ASSENTO", "ENCOSTO": "ENCOSTO",
    "INFANTIL": "INFANTIL", "DESCANSO": "DESCANSO",
}
_SINONIMOS_PRODUTO = {
    "CONF": "CONFORTO", "CONFORT": "CONFORTO", "BERLIN": "BERLIM",
    "NAVE": "NAVEGANTES", "NAVEGANTE": "NAVEGANTES", "NAVEGANTES": "NAVEGANTES",
    "TRAM": "TRAMADO", "TRAMADA": "TRAMADO", "ALU": "ALUMINIO", "ALUM": "ALUMINIO",
    "NAUTICA": "NAUTICO", "CORDAO": "CORDA", "OBI": "OBILONGO", "INF": "INFANTIL",
}
_DETALHES_PRODUTO = {
    "TRICO", "CORDA", "FIBRA", "ARG", "ARGOLA", "TRAMADO", "ALUMINIO", "FERRO",
    "CRA", "ENC", "ARTESANAL", "BAIXA", "BAIXO", "FIO", "COLMEIA", "NAUTICO",
    "PINTADA", "POLIDA", "FLECHA", "OBILONGO", "RET", "QUAD", "COM", "SEM",
    "BRACO", "BRACOS", "ENCAPAGEM", "BASE", "PALITO", "SO", "COMPLETA", "COMPLETO",
}
_STOP_PRODUTO = {"DE", "DO", "DA", "DOS", "DAS", "A", "O", "E", "NA", "NO", "EM"}


def _normalizar_produto_barra(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().upper()
    texto = texto.replace("\\", "/")
    texto = re.sub(r"\b([468])\s*LUG(?:ARES)?\b", r"\1LUG", texto)
    texto = re.sub(r"\b([468])\s*L\b", r"\1LUG", texto)
    texto = re.sub(r"\s*/\s*", "/", texto)
    texto = re.sub(r"[^A-Z0-9/ ]", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def _canon_produto(token):
    return _SINONIMOS_PRODUTO.get(token, token)


def _variantes_produto(texto):
    bruto = _normalizar_produto_barra(texto)
    if not bruto:
        return []
    grupos, achatada = [], []
    for parte in bruto.split():
        alts = [_canon_produto(x) for x in parte.split("/") if x]
        if not alts:
            continue
        achatada.extend(alts)
        aliasavel = len(alts) > 1 and all(
            x not in _TIPOS_PRODUTO and x not in _DETALHES_PRODUTO and x not in _STOP_PRODUTO
            and not re.fullmatch(r"[468]LUG", x) for x in alts
        )
        grupos.append([[x] for x in alts] if aliasavel else [alts])
    variantes = [achatada] if achatada else []
    combos = [[]]
    for escolhas in grupos:
        combos = [prefixo + escolha for prefixo in combos for escolha in escolhas][:64]
    variantes.extend(combos)
    unicas, vistos = [], set()
    for variante in variantes:
        chave = tuple(variante)
        if chave and chave not in vistos:
            vistos.add(chave); unicas.append(variante)
    return unicas


def _partes_produto(tokens):
    tipo = None; resto = []
    for i, tok in enumerate(list(tokens or [])):
        if tok == "BASE":
            continue
        if tipo is None and tok in _TIPOS_PRODUTO and (i <= 1 or not resto):
            tipo = _TIPOS_PRODUTO[tok]; continue
        resto.append(tok)
    identidade, detalhes = [], []
    for tok in resto:
        if tok in _STOP_PRODUTO:
            continue
        if tok in _DETALHES_PRODUTO or re.fullmatch(r"[468]LUG", tok):
            detalhes.append(tok)
        elif tok in _TIPOS_PRODUTO and tipo is None:
            tipo = _TIPOS_PRODUTO[tok]
        else:
            identidade.append(tok)
    return tipo, identidade, detalhes


def _media_fuzzy(origem, alvo):
    if not origem or not alvo: return 0.0
    return sum(max(fuzz.ratio(x, y) for y in alvo) for x in origem) / len(origem)


def _similaridade_identidade(q, a):
    if not q or not a: return 0.0
    sq, sa = set(q), set(a)
    if sq == sa: return 100.0
    qtxt, atxt = " ".join(q), " ".join(a)
    dq, da = _media_fuzzy(q, a), _media_fuzzy(a, q)
    score = 0.42*fuzz.token_set_ratio(qtxt, atxt) + 0.28*dq + 0.12*da + 0.18*fuzz.ratio(qtxt, atxt)
    if sq.issubset(sa) or sa.issubset(sq): score = max(score, 96.0)
    elif sq & sa: score = max(score, 88.0)
    elif max(dq, da) < 78: score = min(score, 58.0)
    return min(100.0, score)


def _similaridade_detalhes(q, a):
    if not q: return 100.0
    if not a: return 25.0
    sq, sa = set(q), set(a)
    if sq == sa: return 100.0
    if sq.issubset(sa): return 98.0
    overlap = len(sq & sa) / len(sq)
    return 0.50*fuzz.token_set_ratio(" ".join(q), " ".join(a)) + 25.0*overlap + 0.25*_media_fuzzy(q, a)


def _pontuar_variante_produto(q, a):
    if not q or not a: return 0.0
    qtxt, atxt = " ".join(q), " ".join(a)
    if qtxt == atxt: return 100.0
    tq, iq, dq = _partes_produto(q); ta, ia, da = _partes_produto(a)
    if not iq or not ia:
        return 0.65*fuzz.token_set_ratio(qtxt, atxt) + 0.35*fuzz.ratio(qtxt, atxt)
    ident = _similaridade_identidade(iq, ia); detalhes = _similaridade_detalhes(dq, da)
    completo = 0.65*fuzz.token_set_ratio(qtxt, atxt) + 0.35*fuzz.ratio(qtxt, atxt)
    if tq and ta and tq != ta:
        return min(38.0, 0.75*ident + 0.15*detalhes + 0.10*completo)
    score = 0.78*ident + 0.14*detalhes + 0.08*completo
    if tq and ta and tq == ta: score += 2.0
    sq, sa = set(iq), set(ia)
    if sq == sa and (not dq or set(dq).issubset(set(da))): score = max(score, 97.0)
    elif (sq.issubset(sa) or sa.issubset(sq)) and (not dq or detalhes >= 80): score = max(score, 92.0)
    if not (sq & sa) and ident < 82: score = min(score, 55.0)
    return max(0.0, min(100.0, score))


def pontuar_nome_produto(nome_consulta, nome_alvo):
    qnorm = _normalizar_produto_barra(nome_consulta).replace("/", " ")
    anorm = _normalizar_produto_barra(nome_alvo).replace("/", " ")
    if qnorm and qnorm == anorm: return 100.0
    qvars, avars = _variantes_produto(nome_consulta), _variantes_produto(nome_alvo)
    if not qvars or not avars: return 0.0
    return max(_pontuar_variante_produto(q, a) for q in qvars for a in avars)


def pontuar(consulta_tokens, item):
    return pontuar_nome_produto(" ".join(consulta_tokens or []), item.get("name") or item.get("norm", ""))


def melhor_match(nome_bruto, catalogo):
    if not _normalizar_produto_barra(nome_bruto) or not catalogo:
        return None, 0.0
    qnorm = _normalizar_produto_barra(nome_bruto).replace("/", " ")
    for item in catalogo:
        if _normalizar_produto_barra(item.get("name") or item.get("norm", "")).replace("/", " ") == qnorm:
            return item, 100.0
    avaliados = sorted(
        ((pontuar_nome_produto(nome_bruto, item.get("name") or item.get("norm", "")), item) for item in catalogo),
        key=lambda x: x[0], reverse=True,
    )
    melhor_score, melhor = avaliados[0]
    if len(avaliados) > 1 and avaliados[1][0] >= 72 and (melhor_score - avaliados[1][0]) < 4.0:
        melhor_score = min(melhor_score, 59.0)
    return melhor, melhor_score


def regra_negocio(nome_bruto):
    return configuracao.aplicar_regras_produto(nome_bruto)


RE_DATA = re.compile(r"^(\d{1,2})/(\d{1,2})$")
RE_FUNCIONARIO_BOLD = re.compile(r"^\*\*(.+?)\*\*$")
RE_PRODUTO = re.compile(r"^(\d+)\s+(.+?)\s*$")
RE_PREFIXO_WHATSAPP = re.compile(r"^\[\d{1,2}:\d{2}(?::\d{2})?,\s*[^\]]+\]\s*[^:]*:\s*")


def parse_texto(texto):
    lancamentos = []
    dia_atual = mes_atual = None
    funcionario_atual = None
    for linha_bruta in texto.splitlines():
        linha = RE_PREFIXO_WHATSAPP.sub("", linha_bruta.strip()).strip()
        if not linha or linha == "---":
            continue
        m = RE_DATA.match(linha)
        if m:
            dia_atual, mes_atual = int(m.group(1)), int(m.group(2))
            funcionario_atual = None
            continue
        m = RE_PRODUTO.match(linha)
        if m and dia_atual is not None:
            lancamentos.append({
                "dia": dia_atual,
                "mes": mes_atual,
                "funcionario": funcionario_atual,
                "produto_bruto": m.group(2).strip(),
                "quantidade": int(m.group(1)),
            })
            continue
        m = RE_FUNCIONARIO_BOLD.match(linha)
        nome = m.group(1).strip() if m else linha.strip("*").strip()
        if nome:
            funcionario_atual = nome
    return lancamentos


def processar_lancamentos(lancamentos, catalogo_cadeiras, catalogo_mesas):
    catalogo = catalogo_cadeiras + catalogo_mesas
    cache = {}
    resultados = []
    for lanc in lancamentos:
        nome_busca, nota = regra_negocio(lanc["produto_bruto"])
        if nome_busca is None:
            resultados.append({**lanc, "match_nome": None, "match_tag": None, "match_row": None,
                               "score": 0.0, "excluido": True, "nota_regra": nota,
                               "excluido_permanente": bool(nota and "permanente" in nota)})
            continue
        if nome_busca == "__SEM_CATALOGO__":
            resultados.append({**lanc, "match_nome": None, "match_tag": None, "match_row": None,
                               "score": 0.0, "excluido": False, "nota_regra": nota,
                               "excluido_permanente": False})
            continue
        chave = normalizar(nome_busca)
        if chave not in cache:
            cache[chave] = melhor_match(nome_busca, catalogo)
        item, score = cache[chave]
        resultados.append({
            **lanc,
            "match_nome": item["name"] if item else None,
            "match_tag": item["tag"] if item else None,
            "match_row": item["row"] if item else None,
            "score": round(score, 1),
            "excluido": False,
            "nota_regra": nota,
            "excluido_permanente": False,
        })
    return resultados


def agregar(resultados):
    aceitar, revisar = _thresholds()
    lancar, revisar_lista, pendentes, excluidos = {}, [], [], []
    for r in resultados:
        if r.get("excluido"):
            excluidos.append(r)
            continue
        if r.get("match_row") is None or r.get("score", 0) < revisar:
            pendentes.append(r)
            continue
        chave = (r["match_tag"], r["match_row"], r["dia"])
        lancar[chave] = lancar.get(chave, 0) + r["quantidade"]
        if r["score"] < aceitar:
            revisar_lista.append(r)
    return lancar, revisar_lista, pendentes, excluidos


def salvar_planilhas(caminho_cadeiras, caminho_mesas, dia_col_cad, dia_col_mesa,
                      lancar, somar_existente, pasta_saida=None):
    def gravar(caminho, dia_col, tag):
        itens = [(chave, qtd) for chave, qtd in lancar.items() if chave[0] == tag]
        if not itens:
            return None
        wb = _abrir_workbook(caminho, read_only=False, data_only=False)
        try:
            ws = wb[SHEET_NAME]
            for (_, row, dia), qtd in itens:
                col = dia_col.get(dia)
                if col is None:
                    continue
                celula = ws.cell(row=row, column=col)
                atual = celula.value
                celula.value = (atual + qtd) if somar_existente and isinstance(atual, (int, float)) else qtd
            destino = tema.caminho_saida_sugerido(caminho, pasta_saida)
            wb.save(destino)
            return destino
        finally:
            wb.close()
    return (
        gravar(caminho_cadeiras, dia_col_cad, "CADEIRAS"),
        gravar(caminho_mesas, dia_col_mesa, "MESAS"),
    )


class SeletorProdutoDialog(tk.Toplevel):
    def __init__(self, parent, catalogo):
        super().__init__(parent)
        self.title("Escolher produto")
        self.geometry("620x500")
        self.minsize(480, 380)
        self.transient(parent)
        self.grab_set()
        self.resultado = None
        self.catalogo = sorted(catalogo, key=lambda x: (x["tag"], x["name"]))
        self.var_busca = tk.StringVar()
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)
        ttk.Label(root, text="Buscar produto", style="Section.TLabel").pack(anchor="w")
        ent = ttk.Entry(root, textvariable=self.var_busca)
        ent.pack(fill="x", pady=(5, 8))
        ent.focus_set()
        self.tree = ttk.Treeview(root, columns=("tipo", "produto"), show="headings", selectmode="browse")
        self.tree.heading("tipo", text="Planilha")
        self.tree.heading("produto", text="Produto")
        self.tree.column("tipo", width=100)
        self.tree.column("produto", width=430)
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<Double-1>", lambda e: self._confirmar())
        self.var_busca.trace_add("write", lambda *_: self._preencher())
        b = ttk.Frame(root)
        b.pack(fill="x", pady=(8, 0))
        ttk.Button(b, text="Cancelar", command=self.destroy).pack(side="right")
        ttk.Button(b, text="Usar produto", style="Accent.TButton", command=self._confirmar).pack(side="right", padx=(0, 8))
        self._preencher()
        self.bind("<Escape>", lambda e: self.destroy())
        self.bind("<Return>", lambda e: self._confirmar())

    def _preencher(self):
        termo = normalizar(self.var_busca.get())
        for i in self.tree.get_children():
            self.tree.delete(i)

        avaliados = []
        for idx, it in enumerate(self.catalogo):
            alvo = f"{it['tag']} {it['name']}"
            n = normalizar(alvo)
            if not termo:
                score = 100.0
            elif termo in n:
                score = 130.0 - max(0, len(n) - len(termo)) * 0.04
            else:
                # A pesquisa manual usa o mesmo matcher do reconhecimento, inclusive '/'.
                score = pontuar_nome_produto(self.var_busca.get(), it["name"])
            if not termo or score >= 42:
                avaliados.append((score, idx, it))

        if termo:
            avaliados.sort(key=lambda x: (-x[0], x[2]["tag"], normalizar(x[2]["name"])))
        else:
            avaliados.sort(key=lambda x: (x[2]["tag"], normalizar(x[2]["name"])))

        # iid usa a posição real no catálogo para _confirmar continuar simples.
        for _score, idx, it in avaliados:
            self.tree.insert("", "end", iid=str(idx), values=(it["tag"], it["name"]))
        if self.tree.get_children():
            self.tree.selection_set(self.tree.get_children()[0])

    def _confirmar(self):
        sel = self.tree.selection()
        if not sel:
            return
        self.resultado = self.catalogo[int(sel[0])]
        self.destroy()


class JanelaExcluidos(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Itens excluídos")
        self.geometry("760x430")
        self.transient(app)
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)
        ttk.Label(root, text="Itens excluídos desta execução", style="Section.TLabel").pack(anchor="w")
        self.tree = ttk.Treeview(root, columns=("data", "produto", "motivo"), show="headings", selectmode="extended")
        for c, t, w in [("data", "Data", 80), ("produto", "Produto recebido", 280), ("motivo", "Motivo", 330)]:
            self.tree.heading(c, text=t); self.tree.column(c, width=w)
        self.tree.pack(fill="both", expand=True, pady=(8, 8))
        for idx, r in enumerate(app.resultados):
            if r.get("excluido"):
                self.tree.insert("", "end", iid=str(idx), values=(f"{r['dia']:02d}/{r['mes']:02d}", r["produto_bruto"], r.get("nota_regra") or "excluído manualmente"))
        ttk.Button(root, text="Restaurar selecionados nesta execução", style="Accent.TButton", command=self._restaurar).pack(anchor="e")

    def _restaurar(self):
        ids = self.tree.selection()
        if not ids:
            return
        for iid in ids:
            self.app._restaurar_indice(int(iid))
        self.app._atualizar_previa()
        self.destroy()


class LancadorProducaoFrame(ttk.Frame):
    def __init__(self, parent):
        super().__init__(parent)
        hist = configuracao.STORE.get("historico", default={}) or {}
        pref_somar = configuracao.STORE.get("preferencias", "somar_existente", default=True)
        self.caminho_cadeiras = tk.StringVar(value=hist.get("cadeiras_arquivo", ""))
        self.caminho_mesas = tk.StringVar(value=hist.get("mesas_arquivo", ""))
        self.somar_existente = tk.BooleanVar(value=bool(pref_somar))
        self.catalogo_cadeiras = []
        self.catalogo_mesas = []
        self.dia_col_cad = {}
        self.dia_col_mesa = {}
        self.resultados = []
        self.lancar = {}
        self._ocupado = False
        self._build()
        self.bind_all("<<ConfiguracaoAlterada>>", self._config_changed, add="+")

    @property
    def catalogo_geral(self):
        return self.catalogo_cadeiras + self.catalogo_mesas

    def _build(self):
        cab = ttk.Frame(self, padding=(16, 12, 16, 8)); cab.pack(fill="x")
        ttk.Label(cab, text="Lançador — Cadeiras / Mesas", style="Header.TLabel").pack(anchor="w")
        ttk.Label(cab, text="Cole o relatório, confira o reconhecimento e grave cópias atualizadas das planilhas.", style="SubHeader.TLabel").pack(anchor="w", pady=(2, 0))

        # Fluxo principal fica sempre no topo, sem ocupar uma faixa grande da tela.
        acoes = ttk.Frame(cab); acoes.pack(fill="x", pady=(9, 0))
        self.btn_processar = ttk.Button(acoes, text="1  ▶  Analisar", style="PrimaryAction.TButton", command=self.processar)
        self.btn_processar.pack(side="left")
        self.btn_salvar = ttk.Button(acoes, text="2  ✓  Lançar nas planilhas", style="LaunchAction.TButton", command=self.salvar, state="disabled")
        self.btn_salvar.pack(side="left", padx=(7, 0))
        self.btn_excluidos = ttk.Button(acoes, text="Excluídos (0)", style="Small.TButton", command=self._abrir_excluidos, state="disabled")
        self.btn_excluidos.pack(side="left", padx=(7, 0))
        ttk.Checkbutton(acoes, text="Somar ao existente", variable=self.somar_existente).pack(side="right")

        self.lbl_status = ttk.Label(self, text="Pronto.", style="Status.TLabel")
        self.lbl_status.pack(side="bottom", fill="x", padx=12, pady=(0, 8))

        rolavel = tema.ScrollableArea(self)
        rolavel.pack(fill="both", expand=True)
        corpo = rolavel.content

        files = ttk.LabelFrame(corpo, text="Planilhas", padding=10); files.pack(fill="x", padx=12, pady=(0, 8))
        files.columnconfigure(0, weight=1)
        self._linha_arquivo(files, 0, "Planilha de Cadeiras", self.caminho_cadeiras, "cadeiras_arquivo")
        self._linha_arquivo(files, 1, "Planilha de Mesas", self.caminho_mesas, "mesas_arquivo")

        paned = ttk.Panedwindow(corpo, orient="vertical"); paned.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        entrada_frame = ttk.Frame(paned); preview = ttk.Frame(paned)
        paned.add(entrada_frame, weight=3); paned.add(preview, weight=7)

        top = ttk.Frame(entrada_frame); top.pack(fill="x", pady=(0, 5))
        ttk.Label(top, text="Mensagem / relatório do WhatsApp", style="Section.TLabel").pack(side="left")
        ttk.Button(top, text="Limpar", style="Small.TButton", command=lambda: self.txt_entrada.delete("1.0", tk.END)).pack(side="right")
        self.txt_entrada = scrolledtext.ScrolledText(entrada_frame, wrap="word", height=9, font=("Consolas", 10)); tema.estilizar_texto(self.txt_entrada); self.txt_entrada.pack(fill="both", expand=True)

        preview.columnconfigure(0, weight=1); preview.rowconfigure(1, weight=1)
        ttk.Label(preview, text="Prévia editável", style="Section.TLabel").grid(row=0, column=0, sticky="w", pady=(3, 5))
        cols = ("data", "func", "recebido", "destino", "planilha", "score", "qtd", "status")
        self.tree = ttk.Treeview(preview, columns=cols, show="headings", selectmode="extended")
        defs = [("data", "Data", 70), ("func", "Funcionário", 130), ("recebido", "Recebido", 220), ("destino", "Produto reconhecido", 230), ("planilha", "Planilha", 85), ("score", "%", 50), ("qtd", "Qtd", 50), ("status", "Status / regra", 250)]
        for c, t, w in defs:
            self.tree.heading(c, text=t); self.tree.column(c, width=w, anchor="w")
        y = ttk.Scrollbar(preview, orient="vertical", command=self.tree.yview); x = ttk.Scrollbar(preview, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)
        self.tree.grid(row=1, column=0, sticky="nsew"); y.grid(row=1, column=1, sticky="ns"); x.grid(row=2, column=0, sticky="ew")
        for tag, (bg, fg) in tema.TAG_CORES.items(): self.tree.tag_configure(tag, background=bg, foreground=fg)
        self.tree.bind("<Double-1>", self._double_click)
        self.tree.bind("<Button-3>", self._menu_contexto)

        self.menu = tk.Menu(self, tearoff=False)
        self.menu.add_command(label="Alterar produto...", command=self._alterar_produto_selecionado)
        self.menu.add_separator()
        self.menu.add_command(label="Excluir nesta execução", command=self._excluir_selecionados)
        self.menu.add_command(label="Excluir sempre este produto", command=self._excluir_sempre_selecionados)
        self.menu.add_command(label="Restaurar nesta execução", command=self._restaurar_selecionados)

    def _linha_arquivo(self, parent, row, rotulo, var, hist_key):
        bloco = ttk.Frame(parent)
        bloco.grid(row=row, column=0, sticky="ew", pady=(0 if row == 0 else 7, 0))
        bloco.columnconfigure(0, weight=1)
        ttk.Label(bloco, text=rotulo, style="Section.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 3))
        ttk.Entry(bloco, textvariable=var).grid(row=1, column=0, sticky="ew")
        ttk.Button(
            bloco, text="Selecionar…", style="Small.TButton",
            command=lambda: self._selecionar_arquivo(var, hist_key, rotulo)
        ).grid(row=1, column=1, padx=(6, 0))

    def _selecionar_arquivo(self, var, hist_key, titulo="Planilha"):
        atual = var.get().strip()
        initial = os.path.dirname(atual) if atual and os.path.isdir(os.path.dirname(atual)) else None
        path = filedialog.askopenfilename(
            title=f"Selecionar {titulo}",
            initialdir=initial,
            filetypes=[("Planilhas Excel", "*.xlsx *.xlsm"), ("Todos os arquivos", "*.*")],
        )
        if path:
            var.set(path)
            configuracao.salvar_historico(hist_key, path)

    def _set_busy(self, busy, texto=None):
        self._ocupado = busy
        self.btn_processar.config(state="disabled" if busy else "normal")
        if busy:
            self.btn_salvar.config(state="disabled")
        if texto:
            self.lbl_status.config(text=texto)

    def _config_changed(self, _e=None):
        if not self.resultados:
            self.somar_existente.set(bool(configuracao.STORE.get("preferencias", "somar_existente", default=True)))

    def processar(self):
        cad, mesa = self.caminho_cadeiras.get().strip(), self.caminho_mesas.get().strip()
        for path, nome in ((cad, "CADEIRAS"), (mesa, "MESAS")):
            ok, erro = tema.validar_caminho_planilha(path)
            if not ok:
                messagebox.showerror(f"Planilha de {nome} inválida", erro); return
        if os.path.abspath(cad) == os.path.abspath(mesa):
            messagebox.showerror("Planilhas iguais", "Selecione arquivos diferentes para CADEIRAS e MESAS."); return
        texto = self.txt_entrada.get("1.0", "end")
        if not texto.strip():
            messagebox.showwarning("Nada colado", "Cole a mensagem antes de processar."); return
        lancamentos = parse_texto(texto)
        if not lancamentos:
            messagebox.showwarning("Nada encontrado", "Não consegui identificar lançamentos no texto."); return

        self._set_busy(True, "Lendo os dois catálogos em segundo plano...")
        def work(report):
            cad_cat, cad_dias = carregar_catalogo(cad, "CADEIRAS"); report((1, 2))
            mesa_cat, mesa_dias = carregar_catalogo(mesa, "MESAS"); report((2, 2))
            return cad_cat, cad_dias, mesa_cat, mesa_dias, processar_lancamentos(lancamentos, cad_cat, mesa_cat)
        def progress(v):
            self.lbl_status.config(text=f"Lendo planilhas... {v[0]}/{v[1]}")
        def done(v):
            self.catalogo_cadeiras, self.dia_col_cad, self.catalogo_mesas, self.dia_col_mesa, self.resultados = v
            configuracao.salvar_historico("cadeiras_arquivo", cad); configuracao.salvar_historico("mesas_arquivo", mesa)
            self._set_busy(False)
            self._atualizar_previa()
        def fail(exc):
            self._set_busy(False, "Falha ao processar."); messagebox.showerror("Erro ao processar", str(exc))
        tema.executar_assincrono(self, work, done, fail, progress)

    def _tag(self, r):
        aceitar, revisar = _thresholds()
        if r.get("excluido"): return "excluido"
        if r.get("match_row") is None or r.get("score", 0) < revisar: return "pendente"
        if r.get("score", 0) < aceitar: return "revisar"
        return "ok"

    def _status(self, r):
        if r.get("excluido"): return r.get("nota_regra") or "Excluído nesta execução"
        if r.get("match_row") is None: return r.get("nota_regra") or "Escolha um produto"
        return r.get("nota_regra") or "OK"

    def _atualizar_previa(self):
        self.lancar, revisar, pendentes, excluidos = agregar(self.resultados)
        for iid in self.tree.get_children(): self.tree.delete(iid)
        for idx, r in enumerate(self.resultados):
            self.tree.insert("", "end", iid=str(idx), tags=(self._tag(r),), values=(
                f"{r['dia']:02d}/{r['mes']:02d}", r.get("funcionario") or "-", r["produto_bruto"],
                r.get("match_nome") or "(não identificado)", r.get("match_tag") or "-", r.get("score", 0),
                r["quantidade"], self._status(r)))
        self.btn_salvar.config(state="normal" if self.lancar and not self._ocupado else "disabled")
        self.btn_excluidos.config(state="normal" if excluidos else "disabled", text=f"Excluídos ({len(excluidos)})")
        self.lbl_status.config(text=f"{len(self.resultados)} linha(s) • {sum(self.lancar.values())} unidade(s) prontas • {len(revisar)} para revisar • {len(pendentes)} pendente(s) • {len(excluidos)} excluída(s)")

    def _selected_indices(self):
        return [int(i) for i in self.tree.selection()]

    def _double_click(self, event):
        iid = self.tree.identify_row(event.y)
        if iid:
            self.tree.selection_set(iid); self._alterar_produto_selecionado()

    def _menu_contexto(self, event):
        iid = self.tree.identify_row(event.y)
        if iid and iid not in self.tree.selection(): self.tree.selection_set(iid)
        if iid:
            try: self.menu.tk_popup(event.x_root, event.y_root)
            finally: self.menu.grab_release()

    def _alterar_produto_selecionado(self):
        ids = self._selected_indices()
        if not ids or not self.catalogo_geral: return
        d = SeletorProdutoDialog(self, self.catalogo_geral); self.wait_window(d)
        if not d.resultado: return
        item = d.resultado
        for idx in ids:
            r = self.resultados[idx]
            r.update(match_nome=item["name"], match_tag=item["tag"], match_row=item["row"], score=100.0,
                     excluido=False, nota_regra="produto definido manualmente", excluido_permanente=False)
        self._atualizar_previa()

    def _excluir_selecionados(self):
        for idx in self._selected_indices():
            self.resultados[idx]["excluido"] = True
            self.resultados[idx]["nota_regra"] = "excluído manualmente nesta execução"
        self._atualizar_previa()

    def _excluir_sempre_selecionados(self):
        ids = self._selected_indices()
        if not ids: return
        nomes = []
        for idx in ids:
            nome = self.resultados[idx]["produto_bruto"]
            if configuracao.adicionar_exclusao(nome, "exato"):
                nomes.append(nome)
            self.resultados[idx]["excluido"] = True
            self.resultados[idx]["excluido_permanente"] = True
            self.resultados[idx]["nota_regra"] = "exclusão permanente (nome exato)"
        self.event_generate("<<ConfiguracaoAlterada>>", when="tail")
        self._atualizar_previa()
        if nomes:
            self.lbl_status.config(text=f"{len(nomes)} produto(s) salvo(s) nas exclusões permanentes.")

    def _restaurar_indice(self, idx):
        if not (0 <= idx < len(self.resultados)): return
        r = self.resultados[idx]
        item, score = melhor_match(r["produto_bruto"], self.catalogo_geral)
        r.update(match_nome=item["name"] if item else None, match_tag=item["tag"] if item else None,
                 match_row=item["row"] if item else None, score=round(score, 1), excluido=False,
                 nota_regra="restaurado manualmente nesta execução", excluido_permanente=False)

    def _restaurar_selecionados(self):
        for idx in self._selected_indices(): self._restaurar_indice(idx)
        self._atualizar_previa()

    def _abrir_excluidos(self):
        JanelaExcluidos(self)

    def salvar(self):
        if not self.lancar:
            return
        total = sum(self.lancar.values())
        confirmar = bool(configuracao.STORE.get("preferencias", "confirmar_salvamento", default=True))
        if confirmar and not messagebox.askyesno("Confirmar lançamento", f"Serão lançadas {total} unidade(s). Deseja continuar?"):
            return
        hist = configuracao.STORE.get("historico", default={}) or {}
        initial = hist.get("pasta_saida", "")
        pasta = filedialog.askdirectory(title="Pasta para salvar as planilhas atualizadas", initialdir=initial if os.path.isdir(initial) else None)
        if not pasta: return
        destinos = []
        if any(k[0] == "CADEIRAS" for k in self.lancar): destinos.append(tema.caminho_saida_sugerido(self.caminho_cadeiras.get(), pasta))
        if any(k[0] == "MESAS" for k in self.lancar): destinos.append(tema.caminho_saida_sugerido(self.caminho_mesas.get(), pasta))
        for d in destinos:
            if not tema.confirmar_sobrescrita(d): return
        self._set_busy(True, "Gravando cópias atualizadas em segundo plano...")
        def work(report):
            return salvar_planilhas(self.caminho_cadeiras.get(), self.caminho_mesas.get(), self.dia_col_cad, self.dia_col_mesa, self.lancar, self.somar_existente.get(), pasta)
        def done(saidas):
            self._set_busy(False)
            self._atualizar_previa()
            configuracao.salvar_historico("pasta_saida", pasta)
            validas = [x for x in saidas if x]
            messagebox.showinfo("Concluído", "Planilha(s) salva(s):\n\n" + "\n".join(validas))
        def fail(exc):
            self._set_busy(False, "Falha ao salvar."); messagebox.showerror("Erro ao salvar", str(exc))
        tema.executar_assincrono(self, work, done, fail)


if __name__ == "__main__":
    root = tk.Tk(); root.title("Lançador de Produção — Cadeiras / Mesas"); root.geometry("1280x820"); root.minsize(980, 650)
    tema.aplicar_tema_escuro(root); LancadorProducaoFrame(root).pack(fill="both", expand=True); root.mainloop()
